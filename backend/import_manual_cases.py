"""Import the supplied manual quotation as historical cases, never active unit prices."""
import argparse
from decimal import Decimal
import hashlib
import json
from pathlib import Path

import openpyxl
from sqlalchemy import select

from app.database import SessionLocal
from app.models import PricingRule, SourceDocument, SourceRecord
from app.pricing_cases import CASE_KIND, arithmetic_formula
from app.services import log_change


def import_cases(db, path):
    path = Path(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    existing = db.scalar(select(SourceDocument).where(SourceDocument.sha256 == digest))
    if existing:
        return existing.report
    wb = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    sheet, cached = wb.active, values.active
    if [sheet[f"{c}2"].value for c in "ABGHKLM"] != [
            "序号", "项目名称", "计量单位", "工程量", "不含税单价", "税率", "含税合价"]:
        raise ValueError("人工报价表表头不匹配；不猜测列含义")
    doc = SourceDocument(filename=path.name, sha256=digest, report={})
    db.add(doc)
    db.flush()
    count, unpriced, subtotal = 0, 0, Decimal(0)
    footer = str(sheet[f"N{sheet.max_row}"].value or "")
    for row in range(4, sheet.max_row):
        item, name, quantity, expression = (sheet[f"{c}{row}"].value for c in ("A", "B", "H", "K"))
        if not isinstance(item, int) or not name:
            continue
        if type(quantity) is not int or quantity <= 0 or not isinstance(expression, str):
            raise ValueError(f"第{row}行数量或公式无效")
        cells = {c.coordinate: str(c.value) for c in sheet[row] if c.value is not None}
        record = SourceRecord(document_id=doc.id, sheet=sheet.title, row_number=row, cells=cells)
        db.add(record)
        db.flush()
        inputs = {f"{c}{row}": sheet[f"{c}{row}"].value for c in ("I", "J")}
        issues = ["公式系数、单价及附加项含义待报价员确认；原式仅用于历史复算"]
        notes = str(sheet[f"N{row}"].value or "")
        if notes:
            issues.append("含原方案调整或补充说明，未经客户确认不得自动替换工艺")
        if "不含税" in footer and cached[f"L{row}"].value:
            issues.append("原表计算含税金额，但末尾备注不含税运，商业口径待确认")
        try:
            amount = arithmetic_formula(expression, inputs)
            if amount <= 0:
                raise ValueError("非正价需核实")
            if abs(amount - Decimal(str(cached[f"K{row}"].value))) > Decimal(".005"):
                raise ValueError("原公式与缓存金额不一致")
            subtotal += amount * quantity
        except ValueError as exc:
            issues.append(str(exc))
            unpriced += 1
        data = {
            "kind": CASE_KIND, "source_item": str(item), "product": str(name),
            "requirements": str(sheet[f"C{row}"].value or ""),
            "dimensions": str(sheet[f"D{row}"].value or ""),
            "quantity": quantity, "unit": str(sheet[f"G{row}"].value),
            "formula": expression, "inputs": inputs,
            "tax_rate": str(cached[f"L{row}"].value),
            "cached_unit_price": str(cached[f"K{row}"].value),
            "cached_taxed_amount": str(cached[f"M{row}"].value),
            "notes": notes, "terms": footer, "issues": issues,
        }
        rule = PricingRule(name=f"人工报价案例 · #{item} {name}"[:255],
                           content=json.dumps(data, ensure_ascii=False), rule_type="case",
                           status="draft", source_record_id=record.id, source_cell=f"K{row}")
        db.add(rule)
        log_change(db, rule, "import_case", "导入客户提供的人工报价，用于复算与核价，不启用为通用单价")
        count += 1
    doc.report = {"kind": CASE_KIND, "cases": count, "unpriced": unpriced,
                  "prices": 0, "rules": count, "source_rows": count, "sheets": [sheet.title],
                  "warnings": ["人工报价历史案例，尚未启用为通用计价模板；原公式系数及商业条件待确认"],
                  "reconciled_untaxed_subtotal": str(subtotal),
                  "cached_taxed_total": str(cached[f"M{sheet.max_row}"].value)}
    log_change(db, doc, "import", "人工报价案例已保留来源、公式、原工程量及待确认条件")
    wb.close()
    values.close()
    return doc.report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    args = parser.parse_args()
    with SessionLocal.begin() as session:
        print(json.dumps(import_cases(session, args.path), ensure_ascii=False))
