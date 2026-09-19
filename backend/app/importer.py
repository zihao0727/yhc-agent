"""Read-only adapters for the two approved price workbooks.

Rows and formula text are retained as evidence. Prices default to active, not verified.
No formula or instruction in a workbook is executed.
"""
import argparse
import hashlib
import json
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from sqlalchemy import select

from .database import SessionLocal
from .models import PriceItem, PricingRule, SourceDocument, SourceRecord
from .services import log_change, product_for


def text(value):
    return "" if value is None else str(value).strip()


def number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
    match = re.fullmatch(r"(\d+(?:\.\d+)?)(?:元)?(?:/平方|/方|/cm|/CM)?(?:起)?", text(value))
    return Decimal(match[1]) if match else None


class Importer:
    def __init__(self, db, path, kind):
        self.db, self.path, self.kind = db, Path(path), kind
        self.sources, self.products = {}, {}
        self.price_count, self.rule_count = 0, 0
        self.signatures = {}
        self.duplicate_count = 0

    def product(self, category, name):
        key = (category, name)
        if key not in self.products:
            self.products[key] = product_for(self.db, category, name)
        return self.products[key]

    def price(self, sheet, row, cell, category, product, **fields):
        item = PriceItem(product_id=self.product(category, product).id,
                         source_record_id=self.sources[sheet, row], source_cell=cell,
                         status="active", review_reason="原表导入，默认启用；计价适用性另行校验", **fields)
        signature = json.dumps([category, product, fields], ensure_ascii=False, sort_keys=True, default=str)
        if signature in self.signatures:
            self.duplicate_count += 1
            item.notes = (item.notes or "") + f"\n原表存在相同配置和单价：{self.signatures[signature]}，保留来源，待确认重复。"
        else:
            self.signatures[signature] = f"{sheet}!{cell}"
        self.db.add(item)
        log_change(self.db, item, "import", "原表导入，默认启用；未改变原始价格性质")
        self.price_count += 1

    def rule(self, sheet, row, cell, content, product=None, rule_type="reference", name=None):
        if not text(content):
            return
        self.db.add(PricingRule(
            product_id=product.id if product else None,
            name=name or text(content).replace("\n", " ")[:90], content=text(content),
            rule_type=rule_type, status="draft",
            source_record_id=self.sources[sheet, row], source_cell=cell,
        ))
        self.rule_count += 1

    def run(self):
        digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        old = self.db.scalar(select(SourceDocument).where(SourceDocument.sha256 == digest))
        if old:
            return {"filename": self.path.name, "skipped": True, "report": old.report}
        self.book = openpyxl.load_workbook(self.path, data_only=False)
        self.doc = SourceDocument(filename=self.path.name, sha256=digest)
        self.db.add(self.doc)
        self.db.flush()
        for sheet in self.book:
            for row in sheet:
                cells = {c.coordinate: c.value for c in row if c.value is not None}
                if not cells:
                    continue
                record = SourceRecord(document_id=self.doc.id, sheet=sheet.title, row_number=row[0].row,
                                      cells=json.loads(json.dumps(cells, ensure_ascii=False, default=str)))
                self.db.add(record)
                self.db.flush()
                self.sources[sheet.title, row[0].row] = record.id
        if self.kind == "a3":
            self.parse_a3()
        elif self.kind == "project":
            self.parse_project()
        else:
            raise ValueError("Unknown workbook type")
        self.doc.report = {"prices": self.price_count, "rules": self.rule_count,
                           "source_rows": len(self.sources), "duplicate_candidates": self.duplicate_count,
                           "sheets": self.book.sheetnames,
                           "warnings": ["全部单价默认启用；计价适用性另行校验，原始公式仅保存，不执行。",
                                        "图片中的产品名称未自动识别；不明确的名称使用待确认占位。",
                                        "价格边界、税费、工艺叠加及重复行需业务审核。"]}
        self.book.close()
        self.db.flush()
        return {"filename": self.path.name, "skipped": False, "report": self.doc.report}

    def parse_project(self):
        required = {"子表1_汇总", "浮雕铝牌", "浮雕铜牌", "材料工艺建议价", "导视报价", "钣金类拆分方案（给客户）"}
        if not required.issubset(self.book.sheetnames):
            raise ValueError("项目报价表结构变化，停止导入")
        # Summary is retained as source evidence, not imported a second time.
        for name in ("浮雕铝牌", "浮雕铜牌"):
            sheet = self.book[name]
            for row in range(2, sheet.max_row + 1):
                v = [sheet.cell(row, c).value for c in range(1, 10)]
                if not v[0] or number(v[7]) is None:
                    continue
                self.price(name, row, f"H{row}", "浮雕牌", text(v[0]), material=text(v[1]),
                           thickness_mm=number(v[2]), spec=text(v[3]), process=text(v[4]),
                           quality=text(v[5]), amount=number(v[7]), unit="m2", price_kind="standard",
                           notes="税费、运费及小面积附加费见规则原文；面积边界待确认。")
        sheet = self.book["材料工艺建议价"]
        for row in range(2, sheet.max_row + 1):
            v = [sheet.cell(row, c).value for c in range(1, 11)]
            if not v[0] or number(v[9]) is None:
                continue
            self.price(sheet.title, row, f"J{row}", "材料与工艺", text(v[0]),
                       material=text(v[1]), thickness_mm=number(v[2]),
                       process=f"{text(v[3])} + {text(v[5])}", amount=number(v[9]), unit="m2",
                       price_kind="suggested", attributes={"material_price": text(v[7]),
                       "process_1": text(v[3]), "process_1_price": text(v[4]),
                       "process_2": text(v[5]), "process_2_price": text(v[6])},
                       notes="原表建议价；保留六位小数，工艺顺序及重复工艺是否有效待确认。")
        for name, start in (("浮雕铝牌", 10), ("材料工艺建议价", 11)):
            for row in self.book[name]:
                for c in row[start - 1:]:
                    if c.value is not None:
                        self.rule(name, c.row, c.coordinate, c.value)
        sheet = self.book["导视报价"]
        for cell in sheet[1]:
            if cell.value is not None:
                self.rule(sheet.title, 1, cell.coordinate, cell.value, rule_type="constraint")
        for row in range(2, sheet.max_row + 1):
            values = {c.column_letter: c.value for c in sheet[row] if c.value is not None}
            if not all(number(values.get(k)) is not None for k in ("B", "D", "F", "H", "J")):
                continue
            cost = sum((number(values[k]) for k in ("B", "D", "F", "H", "J")), Decimal(0))
            for col, margin, complexity in (("K", Decimal(".25"), False), ("L", Decimal(".35"), False),
                                            ("M", Decimal(".25"), True), ("N", Decimal(".35"), True)):
                formula = text(values.get(col))
                base_col = "K" if col == "M" else "L"
                denominator = "0.75" if margin == Decimal(".25") else "0.65"
                expected = (f"={base_col}{row}*1.2" if complexity else
                            f"=ROUND((B{row}+D{row}+F{row}+H{row}+J{row})/{denominator},-2)")
                if formula.replace(" ", "").upper() != expected.upper():
                    raise ValueError(f"未识别公式 {sheet.title}!{col}{row}，停止导入")
                amount = (cost / (1 - margin)).quantize(Decimal("1E2"), rounding=ROUND_HALF_UP)
                if complexity:
                    amount *= Decimal("1.2")
                self.price(sheet.title, row, f"{col}{row}", "导视标识", "导视组合报价",
                           material=text(values.get("A")), process=text(values.get("C")),
                           spec=f"龙骨{text(values.get('E'))} / 灯{text(values.get('G'))} / 人工{text(values.get('I'))}",
                           quality=("复杂类 / " if complexity else "") + f"{int(margin * 100)}%毛利",
                           amount=amount, unit="unknown", price_kind="suggested",
                           attributes={"formula": formula, "components": {k: text(v) for k, v in values.items()},
                                       "gross_margin": str(margin), "complexity_factor": "1.2" if complexity else "1"},
                           notes="按已验证的原表公式独立计算；表头未明确计价单位，计价前必须确认。")
        sheet = self.book["钣金类拆分方案（给客户）"]
        for row in sheet:
            if (sheet.title, row[0].row) not in self.sources:
                continue
            content = "\n".join(f"{c.coordinate}: {c.value}" for c in row if c.value is not None)
            self.rule(sheet.title, row[0].row, f"A{row[0].row}:K{row[0].row}", content,
                      rule_type="case", name=f"钣金参考案例 · 第 {row[0].row} 行")

    def parse_a3(self):
        sheet = self.book["Sheet1"]
        if sheet["D1"].value != "英文-元/CM" or sheet.max_row != 107:
            raise ValueError("A3报价表结构变化，停止导入")
        # Boundaries are explicitly mapped to the inspected original layout.
        blocks = [
            (1, 17, "迷你字", 3, 4, 5, 6), (18, 34, "霓虹灯", 3, 4, 5, 6),
            (35, 47, "平面发光字（包边发光字）", 3, 4, 5, 6),
            (48, 61, "高脚背光字", 3, 4, 5, 6),
            (1, 14, "无边字", 10, 11, 12, 13), (15, 27, "侧发光字", 10, 11, 12, 13),
            (28, 37, "不锈钢平面字", 10, 11, 12, 13),
            (38, 43, "待确认产品（右侧38-43行）", 10, 11, 12, 13),
            (44, 49, "待确认产品（右侧44-49行）", 10, 11, 12, 13),
            (50, 55, "待确认产品（右侧50-55行）", 10, 11, 12, 13),
            (56, 61, "待确认产品（右侧56-61行）", 10, 11, 12, 13),
        ]
        for start, end, name, spec_col, en_col, zh_col, note_col in blocks:
            product = self.product("发光字与金属字", name)
            notes = "\n".join(text(sheet.cell(r, note_col).value) for r in range(start, end + 1)
                              if sheet.cell(r, note_col).value and sheet.cell(r, note_col).value != "备注")
            for row in range(start, end + 1):
                spec = text(sheet.cell(row, spec_col).value)
                has_price = False
                for lang, col in (("en", en_col), ("zh", zh_col)):
                    c = sheet.cell(row, col)
                    value = number(c.value)
                    if value is None:
                        continue
                    has_price = True
                    raw = text(c.value)
                    unit = "m2" if "平方" in raw else "cm"
                    if start in (38, 44, 50, 56) and "以上" in spec:
                        unit = "unknown"
                    self.price(sheet.title, row, c.coordinate, "发光字与金属字", name,
                               spec=spec, language=lang, amount=value, unit=unit,
                               price_kind="starting" if "起" in raw else "standard", notes=notes,
                               attributes={"original_price": raw, "spec_raw": spec})
                if spec and not has_price and spec != "规格":
                    self.rule(sheet.title, row, sheet.cell(row, spec_col).coordinate, spec, product)
                c = sheet.cell(row, note_col)
                if c.value and c.value != "备注":
                    self.rule(sheet.title, row, c.coordinate, c.value, product, "constraint")
        self.a3_flat(sheet, 62, 76, "不锈钢激光切割字（名称待确认）", 3, 4, 5, 6, "201不锈钢", "en")
        self.a3_flat(sheet, 77, 107, "平面铝字", 3, 4, 5, 6, "铝", "en")
        self.a3_flat(sheet, 62, 85, "平面黄铜字（名称待确认）", 10, 11, 12, 13, "黄铜", "en")
        product = self.product("发光字与金属字", "铝板背铣槽发光字（名称待确认）")
        description = ""
        for row in range(87, 98):
            description = text(sheet.cell(row, 9).value) or description
            value = number(sheet.cell(row, 12).value)
            if value is not None:
                thickness = re.match(r"\s*(\d+)mm", description)
                self.price(sheet.title, row, f"L{row}", "发光字与金属字", product.name,
                           material="铝", thickness_mm=Decimal(thickness[1]) if thickness else None,
                           spec=text(sheet.cell(row, 11).value), language="en", unit="cm",
                           amount=value, notes=description)
            if sheet.cell(row, 13).value:
                self.rule(sheet.title, row, f"M{row}", sheet.cell(row, 13).value, product, "constraint")
        for row in range(99, 106):
            self.rule(sheet.title, row, f"H{row}", sheet.cell(row, 8).value,
                      product if row <= 102 else None, "constraint")

    def a3_flat(self, sheet, start, end, name, thickness_col, spec_col, price_col, note_col, material, language):
        product = self.product("发光字与金属字", name)
        thickness = None
        for row in range(start, end + 1):
            raw = text(sheet.cell(row, thickness_col).value)
            match = re.fullmatch(r"(\d+(?:\.\d+)?)(?:mm)?", raw, re.I)
            if match:
                thickness = Decimal(match[1])
            spec = text(sheet.cell(row, spec_col).value)
            value = number(sheet.cell(row, price_col).value)
            if value is not None and spec:
                self.price(sheet.title, row, sheet.cell(row, price_col).coordinate,
                           "发光字与金属字", name, material=material, thickness_mm=thickness,
                           spec=spec, language=language, unit="cm", amount=value,
                           notes="尺寸下限、中文及工艺加价参见产品规则；原表档位边界待确认。")
            for col in (thickness_col, spec_col, note_col):
                content = text(sheet.cell(row, col).value)
                if content and content not in ("备注", "厚度/mm") and (
                        "①" in content or "②" in content or "③" in content or "④" in content
                        or "⑤" in content or "⑥" in content or "计算" in content or col == note_col):
                    self.rule(sheet.title, row, sheet.cell(row, col).coordinate, content, product)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--a3", required=True, type=Path)
    parser.add_argument("--project", required=True, type=Path)
    args = parser.parse_args()
    with SessionLocal.begin() as db:
        reports = [Importer(db, args.a3, "a3").run(), Importer(db, args.project, "project").run()]
    print(json.dumps(reports, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
