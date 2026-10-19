"""Bounded, offline XLSX sources with row-aligned embedded drawings."""
import io
import json
import posixpath
import re
import warnings
import zipfile
import unicodedata
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, coordinate_to_tuple
from PIL import Image, ImageDraw, ImageFont, ImageOps
from .measurement_units import COUNT_ALIASES, ENGINEERING_FACTORS

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ROWS_PER_PAGE = 20
MAX_ROWS = 2000
MAX_COLS = 60
MAX_CELLS = 50000
MAX_EXPANDED = 100 * 1024 * 1024
PREVIEW_IMAGES = 6
IMAGES_PER_BATCH = 6
ROWS_PER_BATCH = 5
IMAGE_BATCH_BYTES = 12 * 1024 * 1024
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "d": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def xml(archive, name):
    data = archive.read(name)
    if b"<!DOCTYPE" in data.replace(b"\x00", b"").upper() or b"<!ENTITY" in data.replace(b"\x00", b"").upper():
        raise ValueError("Excel包含不支持的XML实体")
    return ET.fromstring(data)


def relationships(archive, part, *, allow_missing=False):
    folder, name = posixpath.split(part)
    relpart = posixpath.join(folder, "_rels", name + ".rels")
    if relpart not in archive.namelist():
        return {}
    result = {}
    for rel in xml(archive, relpart):
        if rel.get("TargetMode") == "External":
            continue
        target = rel.get("Target", "")
        resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/")
                                     else posixpath.join(folder, target))
        if resolved.startswith("../"):
            raise ValueError("Excel资源引用无效")
        if resolved not in archive.namelist():
            if allow_missing:
                continue
            raise ValueError("Excel资源引用无效")
        result[rel.get("Id")] = resolved
    return result


def font(size=16):
    for name in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(name).exists():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def render(page, directory, target=None):
    # Text preview and drawings share source-row labels; original images stay separate.
    thumbnails = page["images"][:PREVIEW_IMAGES]
    canvas = Image.new("RGB", (1600, 190 + len(page["rows"]) * 72 + len(thumbnails) * 360), "white")
    draw = ImageDraw.Draw(canvas)
    title_font, body_font = font(24), font(16)
    draw.text((28, 20), f'{page["sheet"]}  {page["row_start"]}-{page["row_end"]}', fill="#25332b", font=title_font)
    y = 75
    for row in page["rows"]:
        text = " | ".join(f'{cell["cell"]}: {cell["value"]}' for cell in row["cells"])
        draw.rectangle((20, y, 1580, y + 68), fill="#f2f6f3" if row["row"] % 2 else "#ffffff")
        draw.text((30, y + 6), text[:100], fill="#26342c", font=body_font)
        if len(text) > 100:
            draw.text((30, y + 32), text[100:200], fill="#26342c", font=body_font)
        y += 72
    if len(page["images"]) > len(thumbnails):
        draw.text((30, y + 4), f'图纸共{len(page["images"])}张；下方为缩略图，全部原图可单独查看',
                  fill="#26342c", font=body_font)
        y += 36
    for item in thumbnails:
        draw.text((30, y + 8), f'图纸锚点 {page["sheet"]}!{item["cell"]}', fill="#26342c", font=body_font)
        with Image.open(directory / item["storage"]) as original:
            thumb = original.copy()
            thumb.thumbnail((1480, 310))
            canvas.paste(thumb, (30, y + 40))
            thumb.close()
        y += 360
    canvas.save(target or directory / f'page-{page["page"]}.png', format="PNG", optimize=True)
    canvas.close()


def prepare_excel(data, directory):
    pages, sources, image_count, cell_count, resource_audit = [], [], 0, 0, []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        names = [member.filename for member in members]
        if len(members) > 3000 or len(names) != len(set(names)):
            raise ValueError("Excel压缩包资源数量异常")
        if sum(member.file_size for member in members) > MAX_EXPANDED:
            raise ValueError("Excel展开内容超过100MB")
        if any(member.flag_bits & 1 or member.file_size > 32 * 1024 * 1024 for member in members):
            raise ValueError("Excel已加密或单项资源过大")
        if any("vbaproject" in name.lower() or "/embeddings/" in name.lower()
               or "/activex/" in name.lower() for name in names):
            raise ValueError("不接受宏、嵌入对象或ActiveX，请另存为普通XLSX")
        if "xl/workbook.xml" not in names:
            raise ValueError("不是有效的XLSX工作簿")
        # Validate every XML part before handing shared strings/styles to openpyxl.
        for name in names:
            if name.lower().endswith((".xml", ".rels")):
                xml(archive, name)
        if any("macroenabled" in node.get("ContentType", "").lower()
               for node in xml(archive, "[Content_Types].xml")):
            raise ValueError("不接受XLSM，请另存为普通XLSX")
        workbook_xml = xml(archive, "xl/workbook.xml")
        wb_rels = relationships(archive, "xl/workbook.xml")
        specs = []
        for sheet in workbook_xml.findall("s:sheets/s:sheet", NS):
            if sheet.get("state", "visible") != "visible":
                continue
            part = wb_rels.get(sheet.get(f'{{{NS["r"]}}}id'))
            if not part:
                raise ValueError("Excel工作表引用无效")
            tree = xml(archive, part)
            cells = tree.findall("s:sheetData/s:row/s:c", NS)
            cell_count += len(cells)
            if cell_count > MAX_CELLS:
                raise ValueError("Excel单元格数量超过50000")
            last_row = 0
            for cell in cells:
                row, col = coordinate_to_tuple(cell.get("r", ""))
                if row > MAX_ROWS or col > MAX_COLS:
                    raise ValueError("Excel行列范围过大，请拆分后上传")
                last_row = max(last_row, row)
            notes = []
            drawings = []
            from .excel_assets import resolve_cell_images
            embedded, audit = resolve_cell_images(archive, cells, directory, image_count)
            drawings.extend(embedded)
            image_count += len(embedded)
            resource_audit.extend({**item, "sheet": sheet.get("name", "")} for item in audit)
            notes.extend(f'图片{item["cell"]}未读取：{item["reason"]}，涉及的文字及尺寸须核实'
                         for item in audit if item["status"] != "resolved")
            sheet_rels = relationships(archive, part)
            for drawing in tree.findall("s:drawing", NS):
                drawing_part = sheet_rels.get(drawing.get(f'{{{NS["r"]}}}id'))
                if not drawing_part:
                    raise ValueError("Excel图纸引用无效")
                drawing_rels = relationships(archive, drawing_part)
                for anchor in xml(archive, drawing_part):
                    origin = anchor.find("d:from", NS)
                    if origin is None:
                        notes.append("存在未绑定单元格的图形，请人工核对所属项目")
                        resource_audit.append({"sheet": sheet.get("name", ""), "cell": "",
                                               "kind": "drawing", "status": "unresolved",
                                               "reason": "unanchored_drawing"})
                        continue
                    row = int(origin.findtext("d:row", "0", NS)) + 1
                    col = int(origin.findtext("d:col", "0", NS)) + 1
                    if row > MAX_ROWS or col > MAX_COLS:
                        raise ValueError("Excel图纸锚点超出限制")
                    blip = anchor.find(".//a:blip", NS)
                    if blip is None:
                        notes.append(f"第{row}行含不支持的形状/图表，请另附图片核对")
                        resource_audit.append({"sheet": sheet.get("name", ""), "row": row,
                            "cell": f"{get_column_letter(col)}{row}", "kind": "drawing",
                            "status": "unresolved", "reason": "unsupported_shape"})
                        continue
                    target = drawing_rels.get(blip.get(f'{{{NS["r"]}}}embed'))
                    if not target:
                        notes.append(f"第{row}行图纸为外部链接，未下载，请另附图片")
                        resource_audit.append({"sheet": sheet.get("name", ""), "row": row,
                            "cell": f"{get_column_letter(col)}{row}", "kind": "drawing",
                            "status": "unresolved", "reason": "external_reference"})
                        continue
                    image_count += 1
                    raw = archive.read(target)
                    with warnings.catch_warnings():
                        warnings.simplefilter("error", Image.DecompressionBombWarning)
                        with Image.open(io.BytesIO(raw)) as image:
                            if image.format not in ("JPEG", "PNG", "WEBP") or getattr(image, "n_frames", 1) != 1:
                                raise ValueError("Excel图纸仅支持单帧JPEG、PNG、WebP")
                            image.load()
                            original = ImageOps.exif_transpose(image).convert("RGB")
                            original.thumbnail((3600, 3600))
                            key = f'excel-image-{image_count}.png'
                            original.save(directory / key, optimize=True)
                            original.close()
                    drawings.append({"row": row, "cell": f"{get_column_letter(col)}{row}", "storage": key})
                    resource_audit.append({"sheet": sheet.get("name", ""), "row": row,
                        "cell": f"{get_column_letter(col)}{row}", "kind": "drawing",
                        "status": "resolved", "storage": key})
                    last_row = max(last_row, row)
            if last_row:
                specs.append((sheet.get("name", "")[:100], tree, drawings, notes, last_row))
        if not specs:
            raise ValueError("Excel没有可读取的可见工作表")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for name, tree, drawings, notes, last_row in specs:
                sheet = workbook[name]
                sheet.reset_dimensions()
                cached = {cell.coordinate: cell.value for row in sheet.iter_rows(
                    min_row=1, max_row=last_row, max_col=MAX_COLS) for cell in row if cell.value is not None}
                indexed = {}
                for cell in tree.findall("s:sheetData/s:row/s:c", NS):
                    address = cell.get("r")
                    row, _ = coordinate_to_tuple(address)
                    value = cached.get(address)
                    formula = cell.findtext("s:f", "", NS)
                    if value is None and not formula:
                        continue
                    if value is not None and len(str(value)) > 4000:
                        raise ValueError("Excel单元格文本过长")
                    item = {"cell": address, "value": "" if value is None else str(value)}
                    if formula:
                        item["formula"] = "=" + formula[:1000]
                        item["formula_status"] = "cached_unverified" if value is not None else "missing_cache"
                    indexed.setdefault(row, []).append(item)
                active = sorted(set(indexed) | {image["row"] for image in drawings})
                # Pack actual rows, not a misleading fixed print area; preserve original indices.
                for start in range(0, len(active), ROWS_PER_PAGE):
                    chunk = active[start:start + ROWS_PER_PAGE]
                    page_drawings = [image for image in drawings if image["row"] in chunk]
                    page = {"page": len(pages) + 1, "sheet": name, "row_start": chunk[0],
                            "row_end": chunk[-1], "rows": [
                                {"row": row, "cells": indexed.get(row, [])} for row in chunk],
                            "images": page_drawings, "warnings": notes,
                            "resource_audit": [item for item in resource_audit if item["sheet"] == name
                                               and (not item.get("row") or item["row"] in chunk)],
                            "header_rows": [{"row": row, "cells": indexed.get(row, [])} for row in active[:3]],
                            "merged_cells": [merge.get("ref") for merge in tree.findall("s:mergeCells/s:mergeCell", NS)]}
                    pages.append(page)
                sources.append(name)
        finally:
            workbook.close()
    manifest = {"format": "xlsx", "parser_version": 2, "sheets": sources, "pages": pages,
                "resource_summary": {"discovered": len(resource_audit),
                    "resolved": sum(item["status"] == "resolved" for item in resource_audit),
                    "unresolved": sum(item["status"] != "resolved" for item in resource_audit)},
                "notice": "公式仅保留原文和未验证缓存值，不执行宏、公式或外部链接；图纸按锚点行关联，跨行图纸归属须复核。"}
    encoded = json.dumps(manifest, ensure_ascii=False)
    if len(encoded) > 600000:
        raise ValueError("Excel原始资料过多，请拆分")
    (directory / "excel.json").write_text(encoded, encoding="utf-8")
    return len(pages)


def source_cells(source):
    return {cell["cell"] for row in [*source["rows"], *source.get("header_rows", [])]
            for cell in row["cells"]} | {image["cell"] for image in source["images"]}


def same_row_range(address, source):
    match = re.fullmatch(r"([A-Z]{1,2}[1-9][0-9]{0,3})\s*[:-]\s*([A-Z]{1,2}[1-9][0-9]{0,3})", address)
    if not match:
        return None
    start, end = coordinate_to_tuple(match[1]), coordinate_to_tuple(match[2])
    cells = source_cells(source)
    if (start[0] != end[0] or start[1] > end[1] or end[1] > MAX_COLS
            or match[1] not in cells or match[2] not in cells):
        return None
    return f"{match[1]}:{match[2]}", [
        cell for cell in sorted(cells, key=coordinate_to_tuple)
        if coordinate_to_tuple(cell)[0] == start[0]
        and start[1] <= coordinate_to_tuple(cell)[1] <= end[1]]


def validate_evidence(evidence, source):
    if evidence.sheet != source["sheet"] or evidence.cell not in source_cells(source):
        raise ValueError("Excel来源工作表或单元格无效")
    if evidence.source_range:
        selected = same_row_range(evidence.source_range, source)
        if selected is None or evidence.cell not in selected[1]:
            raise ValueError("Excel来源范围不存在、跨项目行或不包含引用的单元格")


def project_source_row(line, source):
    columns, header = table_columns(source)
    if not line.source_item or not line.product or not {"item", "product"} <= columns.keys():
        return None
    matches = []
    for row in source["rows"]:
        if row["row"] <= header:
            continue
        cells = {coordinate_to_tuple(cell["cell"])[1]: cell for cell in row["cells"]}
        item, product = cells.get(columns["item"]), cells.get(columns["product"])
        if (item and product and not item.get("formula") and not product.get("formula")
                and item["value"].strip() == line.source_item and product["value"].strip() == line.product):
            matches.append(row)
    return matches[0] if len(matches) == 1 else None


def expand_source_evidence(evidence, source, addresses, original_range=""):
    """Use original cell text, not the model's unverified combined quotation."""
    return [evidence.model_copy(update={"cell": cell["cell"], "text": cell["value"][:1500],
                                      "source_range": original_range})
            for row in source["rows"] for cell in row["cells"]
            if cell["cell"] in addresses and cell["value"].strip() and not cell.get("formula")]


def normalize_evidence(line, source, file_id, page):
    """Fill omitted provenance only when the supplied page uniquely determines it."""
    evidence_items = [*line.evidence, *line.quantity_evidence, *line.dimension_evidence,
                      *(e for m in line.source_measurements for e in m.evidence),
                      *(e for c in line.components for e in c.evidence),
                      *(e for req in line.production_requirements for e in req.evidence),
                      *(e for g in line.letter_groups for e in g.evidence)]
    evidence_items = [e for e in evidence_items if (e.file_id, e.page) == (file_id, page)]
    cells = source_cells(source)
    for evidence in evidence_items:
        cell = evidence.cell
        if "!" in cell:
            sheet, cell = cell.rsplit("!", 1)
            sheet = sheet[1:-1].replace("''", "'") if sheet.startswith("'") and sheet.endswith("'") else sheet
            if sheet != source["sheet"] or evidence.sheet not in ("", sheet):
                continue
        cell = cell.replace("$", "").upper()
        if cell and cell not in cells and same_row_range(cell, source) is None:
            continue
        if evidence.sheet and evidence.sheet != source["sheet"]:
            continue
        evidence.sheet = source["sheet"]
        evidence.cell = cell
    collections = [(line.evidence, 20), *((component.evidence, 10) for component in line.components),
                   *((group.evidence, 10) for group in line.letter_groups),
                   *((req.evidence, 5) for req in line.production_requirements),
                   *((measurement.evidence, 3) for measurement in line.source_measurements)]
    project_row = project_source_row(line, source)
    for collection, limit in collections:
        expanded = []
        for evidence in collection:
            if evidence not in evidence_items or evidence.sheet != source["sheet"] or evidence.source != "text":
                expanded.append(evidence)
                continue
            selected = same_row_range(evidence.cell, source)
            if selected and (project_row is None or coordinate_to_tuple(selected[1][0])[0] == project_row["row"]):
                matches = [cell["cell"] for row in source["rows"] for cell in row["cells"]
                           if cell["cell"] in selected[1] and not cell.get("formula")
                           and evidence.text.strip() and evidence.text.strip() in cell["value"]]
                if len(matches) == 1:
                    replacements = [evidence.model_copy(update={"cell": matches[0], "source_range": selected[0]})]
                elif collection is line.evidence:
                    replacements = expand_source_evidence(evidence, source, selected[1], selected[0])
                else:
                    replacements = []
            elif not evidence.cell and project_row is not None and collection is line.evidence:
                addresses = [cell["cell"] for cell in project_row["cells"]]
                replacements = expand_source_evidence(evidence, source, addresses)
                warning = "模型遗漏单元格地址，已按本批唯一项目序号和名称恢复原文来源，识别字段仍需复核"
                line.uncertainties = list(dict.fromkeys([*line.uncertainties, warning]))[:20]
            else:
                replacements = []
            expanded.extend(replacements or [evidence])
        unique = []
        for evidence in expanded:
            if evidence not in unique:
                unique.append(evidence)
        if len(unique) > limit:
            raise ValueError("Excel来源展开超过字段允许数量，请缩小引用范围")
        collection[:] = unique
    evidence_items = [*line.evidence, *line.quantity_evidence, *line.dimension_evidence,
                      *(e for m in line.source_measurements for e in m.evidence),
                      *(e for c in line.components for e in c.evidence),
                      *(e for req in line.production_requirements for e in req.evidence),
                      *(e for g in line.letter_groups for e in g.evidence)]
    evidence_items = [e for e in evidence_items if (e.file_id, e.page) == (file_id, page)]
    # First use an exact, unique quotation from actual cells, never guessed row numbers.
    for evidence in evidence_items:
        if evidence.cell or evidence.sheet != source["sheet"]:
            continue
        matches = [cell["cell"] for row in source["rows"] for cell in row["cells"]
                   if not cell.get("formula") and cell["value"].strip()
                   and cell["value"].strip() == evidence.text.strip()]
        if len(matches) == 1:
            evidence.cell = matches[0]
    rows = {coordinate_to_tuple(e.cell)[0] for e in line.evidence
            if e in evidence_items and e.sheet == source["sheet"] and e.cell in cells}
    if project_row is not None and rows and rows != {project_row["row"]}:
        raise ValueError("Excel来源行与项目序号和名称不一致")
    if len(rows) == 1:
        row = next(iter(rows))
        columns, _ = table_columns(source)
        row_cells = [cell for item in source["rows"] if item["row"] == row for cell in item["cells"]
                     if cell["value"].strip() and not cell.get("formula")]
        for evidence in evidence_items:
            if evidence.sheet != source["sheet"] or evidence.source != "text":
                continue
            selected = same_row_range(evidence.cell, source)
            if evidence.cell and selected is None:
                continue
            allowed = set(selected[1]) if selected else {cell["cell"] for cell in row_cells}
            matches = [cell["cell"] for cell in row_cells if cell["cell"] in allowed
                       and evidence.text.strip() and evidence.text.strip() in cell["value"]]
            if len(matches) == 1:
                if selected:
                    evidence.source_range = selected[0]
                evidence.cell = matches[0]
        count_cells = [c for item in source["rows"] if item["row"] == row for c in item["cells"]
                       if coordinate_to_tuple(c["cell"])[1] == columns.get("quantity") and not c.get("formula")]
        for evidence in line.quantity_evidence:
            selected = same_row_range(evidence.cell, source)
            if (evidence not in evidence_items or evidence.cell and selected is None
                    or evidence.source != "text" or evidence.sheet != source["sheet"]):
                continue
            for cell in count_cells:
                if selected and cell["cell"] not in selected[1]:
                    continue
                raw = unicodedata.normalize("NFKC", cell["value"]).strip()
                claimed = unicodedata.normalize("NFKC", evidence.text).strip()
                # A known project row disambiguates repeated scalar count cells.
                if raw and (raw == claimed or re.fullmatch(
                        rf"{re.escape(raw)}\s*(?:个|件|套|组|块|只|台|张|樘|面)", claimed)):
                    if selected:
                        evidence.source_range = selected[0]
                    evidence.cell = cell["cell"]
        images = [item["cell"] for item in source["images"] if item["row"] == row]
        for evidence in evidence_items:
            if (not evidence.cell and evidence.sheet == source["sheet"]
                    and evidence.source == "image" and len(set(images)) == 1):
                evidence.cell = images[0]


def table_columns(source):
    plan = source.get("analysis_plan")
    if plan:
        if plan["purpose"] == "price_reference" or plan["layout"] not in ("row_table", "multi_level_header"):
            return {}, 0
        if plan.get("header_row") and "product" in plan["columns"]:
            return {field: coordinate_to_tuple(address)[1] for field, address in plan["columns"].items()}, plan["header_row"]
    aliases = {"item": {"序号", "编号", "项次"}, "product": {"名称", "项目名称", "产品名称", "产品", "品名"},
               "code": {"项目编码", "项目代码", "图号", "图纸编号", "标识编号", "产品代码", "产品编码"},
               "type_note": {"类型", "类型备注", "产品类型"},
               "difficulty": {"难度", "工艺难度", "工艺难度备注", "难度等级"},
               "quantity": {"数量", "制作数量", "工程数量", "工程量", "采购数量", "订购数量",
                            "qty", "quantity"}, "unit": {"单位", "计量单位", "unit"}}
    for row in source.get("header_rows", []):
        mapped = {}
        for cell in row["cells"]:
            value = re.sub(r"\s+", "", cell["value"]).lower()
            for field, names in aliases.items():
                if value in names:
                    mapped[field] = coordinate_to_tuple(cell["cell"])[1]
        if "product" in mapped and len(mapped) >= 2:
            return mapped, row["row"]
    return {}, 0


def table_quantity(evidence, source):
    """Only an actual, non-formula count column can corroborate a table quantity."""
    columns, header = table_columns(source)
    row, col = coordinate_to_tuple(evidence.cell)
    if row <= header or col != columns.get("quantity"):
        return None
    cells = {coordinate_to_tuple(c["cell"])[1]: c
             for item in source["rows"] if item["row"] == row for c in item["cells"]}
    cell = cells.get(col)
    if not cell or cell.get("formula"):
        return None
    unit_cell = cells.get(columns.get("unit"), {})
    if unit_cell.get("formula"):
        return None
    unit = unicodedata.normalize("NFKC", unit_cell.get("value", "")).strip().lower()
    units = COUNT_ALIASES
    if unit and unit not in units:
        return None
    value = unicodedata.normalize("NFKC", cell["value"]).strip().lower()
    match = re.fullmatch(r"([0-9]+)(?:\.0+)?\s*([A-Za-z\u4e00-\u9fff]*)", value)
    if (not match or match[2] and match[2] not in units
            or unit and match[2] and units[unit] != units[match[2]]):
        return None
    count = int(match[1])
    return count if 0 < count <= 1000000 else None


def table_measurement(source, row_number, file_id, page):
    """Ordered length/area is source engineering quantity, never a fabricated piece count."""
    from .requirements_schemas import Evidence, SourceMeasurement
    columns, header = table_columns(source)
    if row_number <= header or "quantity" not in columns:
        return None
    cells = {coordinate_to_tuple(c["cell"])[1]: c
             for row in source["rows"] if row["row"] == row_number for c in row["cells"]}
    quantity = cells.get(columns["quantity"])
    unit_cell = cells.get(columns.get("unit"))
    if not quantity or quantity.get("formula") or (unit_cell and unit_cell.get("formula")):
        return None
    value = unicodedata.normalize("NFKC", quantity["value"]).strip()
    unit = unicodedata.normalize("NFKC", unit_cell["value"]).strip().lower() if unit_cell else ""
    match = re.fullmatch(r"([0-9]+(?:\.[0-9]+)?)\s*([A-Za-z0-9²³㎡\u4e00-\u9fff^]*)", value)
    if not match:
        return None
    factors = ENGINEERING_FACTORS
    suffix = match[2].lower()
    if suffix and unit and factors.get(suffix) != factors.get(unit):
        return None
    kind = factors.get(unit or suffix)
    amount = Decimal(match[1])
    if not kind or amount <= 0:
        return None
    evidence = [Evidence(file_id=file_id, page=page, sheet=source["sheet"],
                         cell=c["cell"], text=c["value"][:1500])
                for c in [quantity, unit_cell] if c]
    from pydantic import ValidationError
    try:
        return SourceMeasurement(value=amount * Decimal(kind[1]), unit=kind[0], evidence=evidence)
    except ValidationError:
        return None


def financial_summary(name):
    return bool(re.fullmatch(
        r"\s*(?:(?:税费|税点|税金|税额)(?:\s*[\d.]+\s*[%％])?|(?:包装|打包)?运费|包装费|运输费)\s*",
        name))


def validate_excel_rows(result, sources):
    """Only recognizable customer-table columns supply conservative row fallbacks."""
    from .requirements_schemas import ExtractedLine, Evidence, QuantityEvidence
    from .source_coverage import account_rows, row_kind
    result = account_rows(result, sources)
    for (file_id, page), source in sources.items():
        if source.get("analysis_plan", {}).get("purpose") == "price_reference":
            result.lines = [line for line in result.lines if line.text_evidence or not line.evidence
                            or any((e.file_id, e.page) != (file_id, page) for e in line.evidence)]
            result.questions = list(dict.fromkeys([*result.questions,
                f'Excel {source["sheet"]}为价格参考资料，不从价目表自动生成客户采购数量']))[:30]
            continue
        columns, header_row = table_columns(source)
        if not columns:
            warning = f'Excel {source["sheet"]} 第{page}页未找到明确的客户明细表头，逐行覆盖须人工核实'
            result.questions = list(dict.fromkeys([*result.questions, warning]))[:30]
            continue
        for row in source["rows"]:
            if row["row"] <= header_row:
                continue
            if row_kind(row, header_row, columns) == "non_product":
                continue
            cells = {coordinate_to_tuple(cell["cell"])[1]: cell for cell in row["cells"]}
            name_cell = cells.get(columns["product"])
            if not name_cell or not name_cell["value"].strip() or name_cell.get("formula"):
                continue
            name = name_cell["value"].strip()
            if re.match(r"^(?:小计|合计|总计|备注|说明|注意)(?:[:：\s]|$)", name):
                continue
            item_cell = cells.get(columns.get("item"))
            item = item_cell["value"] if item_cell else ""
            associated = [line for line in result.lines if any(
                evidence.file_id == file_id and evidence.page == page and evidence.sheet == source["sheet"]
                and coordinate_to_tuple(evidence.cell)[0] == row["row"] for evidence in line.evidence)]
            if financial_summary(name):
                # A financial summary in a product column is not a new fabricated product.
                result.lines = [line for line in result.lines if not any(line is item for item in associated)]
                result.questions = list(dict.fromkeys([*result.questions,
                    f'原表{source["sheet"]}!{name_cell["cell"]}为费用摘要“{name}”，保留原资料，不作为产品明细计价']))[:30]
                continue
            if not associated:
                raw = " | ".join(f'{cell["cell"]}: {cell["value"]}' for cell in row["cells"])
                line = ExtractedLine(source_item=item, product=name[:150], notes=raw[:4000],
                    evidence=[Evidence(file_id=file_id, page=page, sheet=source["sheet"],
                                       cell=name_cell["cell"], text=raw[:1500])],
                    uncertainties=["该行由Excel单元格补回，模型未完整提取，请核对图纸、规格及数量"])
                result.lines.append(line)
                associated = [line]
            for line in associated:
                measurement = table_measurement(source, row["row"], file_id, page)
                table_claims = [m for m in line.source_measurements if any(
                    e.file_id == file_id and e.page == page and e.sheet == source["sheet"]
                    and coordinate_to_tuple(e.cell)[0] == row["row"] and e.source == "text"
                    and not any(image["cell"] == e.cell for image in source["images"])
                    for e in m.evidence if e.cell)]
                line.source_measurements = [m for m in line.source_measurements if m not in table_claims]
                if table_claims and measurement is None:
                    line.uncertainties = list(dict.fromkeys([*line.uncertainties,
                        "模型工程量输出不能由Excel数量及单位栏证明，未采用，请核对原资料"]))[:20]
                if measurement is not None:
                    line.source_measurements = [m for m in line.source_measurements if not any(
                        e.file_id == file_id and e.page == page and e.sheet == source["sheet"]
                        and coordinate_to_tuple(e.cell)[0] == row["row"] for e in m.evidence if e.cell)]
                    line.source_measurements = [*line.source_measurements, measurement][:10]
                    if not any(e.validation_status == "verified" for e in line.quantity_evidence):
                        line.quantity = None
                    warning = "Excel数量栏为长度/面积工程量，已保留原值及单位，不作为成品件数"
                    line.uncertainties = list(dict.fromkeys([*line.uncertainties, warning]))[:20]
                count_cell = cells.get(columns.get("quantity"))
                if count_cell:
                    quantity_source = QuantityEvidence(file_id=file_id, page=page, sheet=source["sheet"],
                        cell=count_cell["cell"], text=count_cell["value"][:1500], quantity=1,
                        validation_status="verified")
                    count = table_quantity(quantity_source, source)
                    conflict = any("quantity不一致" in warning or ("数量" in warning and "冲突" in warning)
                                   for warning in line.uncertainties)
                    if count is not None and not conflict:
                        quantity_source.quantity = count
                        unit_cell = cells.get(columns.get("unit"), {})
                        unit_label = unicodedata.normalize("NFKC", unit_cell.get("value", "")).strip().lower()
                        line.quantity_unit = (unit_label if re.search(r"[\u4e00-\u9fff]", unit_label)
                                              else {"piece": "件", "set": "套"}.get(COUNT_ALIASES.get(unit_label), ""))
                        if line.quantity not in (None, count) or any(
                                evidence.quantity != count for evidence in line.quantity_evidence
                                if evidence.validation_status != "rejected"):
                            line.quantity = None
                            line.uncertainties = list(dict.fromkeys([*line.uncertainties,
                                "数量识别与Excel数量栏存在冲突，已保留原文，须确认"]))[:20]
                        else:
                            line.quantity = count
                        if quantity_source not in line.quantity_evidence:
                            line.quantity_evidence = [*line.quantity_evidence, quantity_source][-20:]
                    elif count is None and not any(
                            evidence.validation_status == "verified" for evidence in line.quantity_evidence):
                        line.quantity = None
                        warning = "Excel数量栏尚不能证明成品件数，请核对计量单位、公式缓存或小数工程量"
                        line.uncertainties = list(dict.fromkeys([*line.uncertainties, warning]))[:20]
                for field in ("type_note", "difficulty"):
                    cell = cells.get(columns.get(field))
                    if not cell or not cell["value"] or cell.get("formula"):
                        continue
                    value = cell["value"]
                    if field == "type_note":
                        if line.type_note and line.type_note != value:
                            line.uncertainties = list(dict.fromkeys([*line.uncertainties,
                                "类型备注与Excel原文不同，已保留单元格原值"]))[:20]
                        line.type_note = value[:255]
                    else:
                        line.difficulty_note = value[:500]
                        grade = re.fullmatch(r"\s*([ABCabc])(?:\s*类|\s*级)?\s*", value)
                        if grade:
                            line.difficulty = grade[1].upper()
                        else:
                            line.difficulty = "unknown"
                            line.uncertainties = list(dict.fromkeys([*line.uncertainties,
                                f"原始工艺难度备注：{value[:350]}，等级待核实"]))[:20]
                    if not any(e.cell == cell["cell"] and e.sheet == source["sheet"] and e.file_id == file_id
                               and e.page == page for e in line.evidence):
                        line.evidence = [*line.evidence, Evidence(file_id=file_id, page=page,
                            sheet=source["sheet"], cell=cell["cell"], text=value[:1500])][-20:]
    if len(result.lines) > 100:
        raise ValueError("Excel需求超过100行，请拆分会话")
    return account_rows(result, sources)


def page_source(directory, page):
    from .excel_manifest import load_manifest
    manifest = load_manifest(directory, upgrade=True)
    source = manifest["pages"][page - 1]
    return {**source, "notice": manifest["notice"], "parser_version": manifest.get("parser_version", 1),
            "upgrade_status": manifest.get("upgrade_status", "")}


def row_source(source, rows):
    selected = set(rows)
    _, header = table_columns(source)
    return {**source, "rows": [row for row in source["rows"] if row["row"] in selected],
            "header_rows": [row for row in source.get("header_rows", []) if row["row"] <= header],
            "images": [image for image in source["images"] if image["row"] in selected],
            "row_start": min(rows), "row_end": max(rows)}


def detail_images(directory, page, selection=None, rows=None):
    source = page_source(directory, page)
    if rows is not None:
        source = row_source(source, rows)
    images = source["images"] if selection is None else source["images"][selection[0]:selection[1]]
    for item in images:
        with Image.open(directory / item["storage"]) as original:
            image = original.copy()
            image.thumbnail((1800, 1800))
            output = io.BytesIO()
            image.save(output, "PNG")
            image.close()
            yield (f'工作表{source["sheet"]}!{item["cell"]}锚点图纸 {item["storage"]}，须核对对应项目',
                   output.getvalue())


def image_batches(directory, page, rows=None):
    """Bound request memory and bytes, not the customer's total drawing count."""
    start, count, byte_size = 0, 0, 0
    for index, (_, data) in enumerate(detail_images(directory, page, rows=rows)):
        if count and (count >= IMAGES_PER_BATCH or byte_size + len(data) > IMAGE_BATCH_BYTES):
            yield (start, index)
            start, count, byte_size = index, 0, 0
        count += 1
        byte_size += len(data)
    if count:
        yield (start, start + count)


def merge_drawing_results(results):
    """Repeated views refine one item; differing values remain review questions."""
    from .requirements_schemas import ExtractedLine, ExtractionResult
    defaults = ExtractedLine().model_dump()
    rows, conflicts = {}, {}
    lists = {"evidence", "text_evidence", "quantity_evidence", "dimension_evidence", "uncertainties",
             "source_measurements"}
    summaries, questions = [], []
    for result in results:
        summaries.append(result.summary)
        questions.extend(result.questions)
        for incoming in result.lines:
            source_rows = tuple(sorted({coordinate_to_tuple(e.cell)[0] for e in incoming.evidence if e.cell}))
            key = (incoming.source_item or source_rows, incoming.product)
            if key not in rows:
                rows[key] = incoming.model_dump()
                conflicts[key] = set()
                continue
            merged = rows[key]
            for field, value in incoming.model_dump().items():
                if field in lists:
                    for item in value:
                        if item not in merged[field]:
                            merged[field].append(item)
                    merged[field] = merged[field][:10 if field == "source_measurements" else 20]
                elif field == "notes":
                    if value and value not in merged[field]:
                        merged[field] = (merged[field] + "\n" + value).strip()[:4000]
                elif value != defaults[field] and field not in conflicts[key]:
                    if merged[field] == defaults[field]:
                        merged[field] = value
                    elif merged[field] != value:
                        conflicts[key].add(field)
                        merged[field] = defaults[field]
            for field in sorted(conflicts[key]):
                warning = f"同项目分批图纸的{field}不一致，待核实；重复视图未累加数量"
                merged["uncertainties"] = list(dict.fromkeys([*merged["uncertainties"], warning]))[:20]
    return ExtractionResult(lines=[ExtractedLine.model_validate(row) for row in rows.values()],
        summary="\n".join(dict.fromkeys(summaries))[:4000],
        questions=list(dict.fromkeys(questions))[:30])
