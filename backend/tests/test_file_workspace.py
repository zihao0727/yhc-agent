import copy
import io
import json
import os
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from app import file_sandbox, file_workspace as workspace
from app.config import settings
from app.customer_files import prepare_file, pdf_page_text
from app.file_workspace import FileScriptInput, ArtifactInput


def source_file(quantity=2, file_id=1, job_id=7, extra_row=False, unit=None):
    book = Workbook()
    book.active.title = "清单"
    book.active.append(["序号", "产品名称", "数量", "计量单位"])
    book.active.append([1, "标识牌", quantity, unit])
    if extra_row:
        book.active.append([2, "指引牌", 3])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return SimpleNamespace(id=file_id, job_id=job_id, **prepare_file(stream.getvalue(), "客户.xlsx"))


def extracted(file, quantity=2):
    return {"extraction": {"lines": [{
        "source_item": "1", "product": "标识牌", "quantity": quantity,
        "evidence": [{"file_id": file.id, "page": 1, "sheet": "清单", "cell": "B2", "text": "标识牌"}],
        "quantity_evidence": [{"file_id": file.id, "page": 1, "sheet": "清单", "cell": "C2",
                               "text": str(quantity), "quantity": quantity, "raw_quantity": str(quantity)}],
    }]}}


def args(file, **overrides):
    return FileScriptInput(**{"file_ids": [file.id], "script": "print('parse')", "purpose": "核对原表", **overrides})


def test_workspace_copies_only_selected_original_and_archives_result(monkeypatch):
    file = source_file()
    output = extracted(file)
    def run(directory, check):
        manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
        assert manifest["files"][0]["path"] == "/input/file-1.xlsx"
        assert sorted(p.name for p in directory.iterdir()) == ["file-1.xlsx", "manifest.json", "script.py"]
        assert manifest["source_pages"][0]["file_id"] == file.id
        return {"ok": True, "output": output}
    monkeypatch.setattr(file_sandbox, "run_python", run)
    result = workspace.run_script(7, [file], args(file))
    assert result["status"] == "validated"
    assert result["validation"]["extraction"]["lines"][0]["quantity"] == 2
    assert not list((workspace.job_root(7) / result["artifact_id"] / "input").glob("file-*"))
    assert workspace.read_artifact(7, result["artifact_id"])["script_sha256"]
    assert workspace.artifact_page(7, ArtifactInput(
        artifact_id=result["artifact_id"], section="script"))["text"] == args(file).script
    with pytest.raises(ValueError, match="不属于"):
        workspace.read_artifact(8, result["artifact_id"])
    with pytest.raises(ValueError):
        workspace.read_artifact(7, "../secrets")


def test_failed_script_can_be_repaired_and_failure_is_retained(monkeypatch):
    file = source_file()
    answers = iter([{"ok": False, "error": "KeyError: column"},
                    {"ok": True, "output": extracted(file)}])
    monkeypatch.setattr(file_sandbox, "run_python", lambda *a: next(answers))
    failed = workspace.run_script(7, [file], args(file))
    fixed = workspace.run_script(7, [file], args(file, script="print('fixed')",
                                                 previous_artifact=failed["artifact_id"]))
    assert failed["status"] == "execution_failed"
    assert fixed["status"] == "validated"
    assert fixed["previous_artifact"] == failed["artifact_id"]
    assert workspace.read_artifact(7, failed["artifact_id"])["execution"]["error"] == "KeyError: column"


@pytest.mark.parametrize("change", ["foreign_file", "wrong_cell", "wrong_quantity", "duplicate", "no_evidence", "wrong_evidence_shape"])
def test_untrusted_script_cannot_declare_invalid_extraction_valid(monkeypatch, change):
    file = source_file()
    output = extracted(file)
    line = output["extraction"]["lines"][0]
    if change == "foreign_file":
        line["evidence"][0]["file_id"] = 999
    elif change == "wrong_cell":
        line["evidence"][0]["cell"] = "Z999"
    elif change == "wrong_quantity":
        line["quantity"] = 99
        line["quantity_evidence"][0].update(quantity=99, text="99", raw_quantity="99")
    elif change == "duplicate":
        output["extraction"]["lines"].append(copy.deepcopy(line))
    elif change == "wrong_evidence_shape":
        line["evidence"] = line["evidence"][0]
    else:
        line["evidence"] = []
    output["checks"] = [{"passed": True}]
    monkeypatch.setattr(file_sandbox, "run_python", lambda *a: {"ok": True, "output": output})
    result = workspace.run_script(7, [file], args(file))
    assert not result["validation"]["valid"]
    with pytest.raises(ValueError, match="独立校验"):
        workspace.save_adapter(7, result["artifact_id"], "无效适配器")


def test_missing_count_is_recovered_only_from_original():
    file = source_file()
    texts = {(file.id, 1): pdf_page_text(file, 1)}
    output = extracted(file)
    output["extraction"]["lines"][0]["quantity"] = None
    output["extraction"]["lines"][0]["quantity_evidence"] = []
    result = workspace.validate_output(output, [file], texts)
    # The native reader may recover a quantity; it must be the actual source value.
    assert result["extraction"]["lines"][0]["quantity"] in (None, 2)
    if result["extraction"]["lines"][0]["quantity"] is None:
        assert not result["valid"] and result["quantity_missing"] == [0]


def test_omitted_product_is_restored_and_flagged():
    file = source_file(extra_row=True)
    result = workspace.validate_output(extracted(file), [file], {(1, 1): pdf_page_text(file, 1)})
    assert not result["valid"]
    assert result["restored_rows"] == 1
    assert len(result["extraction"]["lines"]) == 2
    assert result["accepted"]


def test_renaming_same_source_does_not_hide_duplicates():
    file = source_file()
    output = extracted(file)
    another = copy.deepcopy(output["extraction"]["lines"][0])
    another.update(product="别名", source_item="2")
    output["extraction"]["lines"].append(another)
    result = workspace.validate_output(output, [file], {(1, 1): pdf_page_text(file, 1)})
    assert not result["valid"] and not result["accepted"]
    assert result["duplicate_sources"] == [[0, 1]]


def test_wrong_unit_is_preserved_as_engineering_quantity():
    file = source_file(unit="m")
    text = json.loads(pdf_page_text(file, 1))
    output = extracted(file)
    result = workspace.validate_output(output, [file], {(1, 1): json.dumps(text, ensure_ascii=False)})
    assert not result["valid"]
    assert result["extraction"]["lines"][0]["quantity"] is None
    assert result["extraction"]["lines"][0]["source_measurements"][0]["unit"] == "m"


def test_candidate_adapter_needs_new_source_and_always_revalidates(monkeypatch):
    file = source_file()
    monkeypatch.setattr(file_sandbox, "run_python", lambda *a: {"ok": True, "output": extracted(file)})
    first = workspace.run_script(7, [file], args(file))
    candidate = workspace.save_adapter(7, first["artifact_id"], "三列表格")
    assert candidate["status"] == "candidate"
    assert workspace.save_adapter(7, first["artifact_id"], "重复保存")["status"] == "candidate"
    other = source_file(3, file_id=2, job_id=8)
    monkeypatch.setattr(file_sandbox, "run_python", lambda *a: {"ok": True, "output": extracted(other, 3)})
    reused = workspace.run_adapter(8, [other], workspace.AdapterInput(
        adapter_id=candidate["adapter_id"], file_ids=[2]))
    assert reused["validation"]["extraction"]["lines"][0]["quantity"] == 3
    assert workspace.list_adapters()[0]["status"] == "verified"
    monkeypatch.setattr(file_sandbox, "run_python", lambda *a: {"ok": True, "output": extracted(other, 999)})
    invalid = workspace.run_adapter(8, [other], workspace.AdapterInput(
        adapter_id=candidate["adapter_id"], file_ids=[2]))
    assert not invalid["validation"]["valid"]


def test_original_checksum_change_stops_execution(monkeypatch):
    file = source_file()
    file.sha256 = "0" * 64
    def never(*args):
        pytest.fail("Changed source must not execute")
    monkeypatch.setattr(file_sandbox, "run_python", never)
    with pytest.raises(ValueError, match="校验和"):
        workspace.run_script(7, [file], args(file))


def test_environment_unavailable_never_executes_host_python(monkeypatch, tmp_path):
    monkeypatch.setattr(file_sandbox, "availability", lambda: {"available": False, "reason": "offline"})
    monkeypatch.setattr(file_sandbox.subprocess, "Popen", lambda *a, **k: pytest.fail("must not execute"))
    result = file_sandbox.run_python(tmp_path)
    assert result["cause"] == "service_unavailable"


def test_recovery_script_uses_existing_extraction_validator(monkeypatch):
    from app import deepseek
    from app.extraction_recovery import recover_extraction
    from test_agent_loop import call
    from test_extraction_recovery import checkpoint
    file = source_file()
    monkeypatch.setattr(file_sandbox, "run_python", lambda *a: {"ok": True, "output": extracted(file)})
    monkeypatch.setattr(deepseek, "agent_turn", lambda *a: (
        call("run_file_script", **args(file).model_dump()), {"total_tokens": 7}))
    cache, _, _, events = checkpoint()
    result, attempts, usage = recover_extraction(
        {"lines": []}, {"stage": "schema"}, [file], {(1, 1): pdf_page_text(file, 1)},
        [], "", "", cache, "batch")
    assert result.lines[0].quantity == 2
    assert attempts == 1 and usage["total_tokens"] == 7
    assert workspace.list_adapters()[0]["status"] == "candidate"


def test_api_is_order_scoped(client, test_db):
    from app.models import RequirementJob
    with test_db.begin() as db:
        job = RequirementJob(title="Workspace", requirements=[])
        db.add(job)
        db.flush()
        job_id = job.id
    response = client.get(f"/api/jobs/{job_id}/file-workspace/bad")
    assert response.status_code == 404
    assert client.get("/api/jobs/999999/file-workspace").status_code == 404


SCRIPT = '''
import json
from openpyxl import load_workbook
m = json.load(open("/input/manifest.json", encoding="utf-8"))
f = m["files"][0]
w = load_workbook(f["path"], data_only=True)
s = w.active
lines = []
for row in s.iter_rows(min_row=2, max_col=3):
    number, product, quantity = row
    lines.append({"source_item": str(number.value), "product": product.value, "quantity": quantity.value,
      "evidence": [{"file_id": f["file_id"], "page": 1, "sheet": s.title, "cell": product.coordinate, "text": product.value}],
      "quantity_evidence": [{"file_id": f["file_id"], "page": 1, "sheet": s.title, "cell": quantity.coordinate,
        "text": str(quantity.value), "quantity": quantity.value, "raw_quantity": str(quantity.value)}]})
json.dump({"extraction": {"lines": lines}}, open("/work/result.json", "w", encoding="utf-8"), ensure_ascii=False)
'''


@pytest.mark.skipif(os.environ.get("FIREFLY_SANDBOX_INTEGRATION") != "1", reason="Explicit Docker integration run")
def test_real_container_reads_workbook_and_enforces_boundaries(monkeypatch):
    assert file_sandbox.availability()["available"]
    file = source_file()
    sandbox_checks = '''
import os, socket
assert not os.environ.get("DEEPSEEK_API_KEY")
assert not os.path.exists("/input/.env")
try:
    open("/input/manifest.json", "w")
    raise AssertionError("input was writable")
except OSError:
    pass
try:
    open("/etc/host-change", "w")
    raise AssertionError("root was writable")
except OSError:
    pass
s = socket.socket()
s.settimeout(1)
assert s.connect_ex(("1.1.1.1", 443)) != 0
s.close()
'''
    result = workspace.run_script(7, [file], args(file, script=sandbox_checks + SCRIPT))
    assert result["status"] == "validated", result
    assert result["validation"]["extraction"]["lines"][0]["quantity"] == 2
    monkeypatch.setattr(settings(), "file_workspace_timeout", 2)
    failed = workspace.run_script(7, [file], args(file, script="while True: pass"))
    assert failed["status"] == "execution_failed"
    assert "超时" in failed["execution"]["error"]
