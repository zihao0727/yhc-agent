import io
import json
import zipfile
from types import SimpleNamespace
from xml.etree import ElementTree as ET

import pytest
import httpx
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from app.customer_excel import MEDIA_TYPE, page_source, validate_excel_rows, image_batches, merge_drawing_results
from app.customer_files import FileRejected, cleanup_file, prepare_file, storage_path
from app.excel_manifest import ensure_preview
from app.requirements_schemas import ExtractionResult, ExtractedLine, Evidence
from app.config import settings
from app import deepseek
from test_deepseek import configure_transport


def workbook_bytes(rows=3, drawing=True, image_count=1):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客户明细"
    sheet.append(["序号", "产品名称", "类型备注", "工艺难度备注", "数量", "费用公式"])
    for index in range(rows):
        sheet.append([index + 1, f"项目{index + 1}", "发光字", "B", 2, "=1+2"])
    for index in range(image_count if drawing else 0):
        stream = io.BytesIO()
        Image.new("RGB", (300, 200), (index % 256, 170, 128)).save(stream, "PNG")
        stream.seek(0)
        sheet.add_image(ExcelImage(stream), "G2")
    other = workbook.create_sheet("造型")
    other.append(["序号", "产品名称", "类型备注"])
    other.append([1, "立牌", "双面"])
    hidden = workbook.create_sheet("隐藏")
    hidden.append(["不发送给模型"])
    hidden.sheet_state = "hidden"
    result = io.BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


def modify_zip(data, name, transform):
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(result, "w") as output:
        for part in original.infolist():
            value = original.read(part.filename)
            output.writestr(part.filename, transform(value) if part.filename == name else value)
    return result.getvalue()


def test_excel_pages_images_formulas_and_hidden_sheets():
    file = prepare_file(workbook_bytes(rows=22), "客户.xlsx")
    directory = storage_path(file["storage_key"])
    assert file["media_type"] == MEDIA_TYPE and file["page_count"] == 3
    first = page_source(directory, 1)
    assert first["sheet"] == "客户明细" and first["row_end"] == 20
    assert first["images"][0]["cell"] == "G2"
    formula = first["rows"][1]["cells"][-1]
    assert formula["formula"] == "=1+2" and formula["formula_status"] == "missing_cache"
    assert formula["value"] == ""
    second = page_source(directory, 2)
    assert second["rows"][0]["row"] == 21 and second["header_rows"][0]["row"] == 1
    assert page_source(directory, 3)["sheet"] == "造型"
    with Image.open(directory / first["images"][0]["storage"]) as image:
        assert image.size == (300, 200)
    cleanup_file(file["storage_key"])
    assert not directory.exists()


def test_excel_row_coverage_and_authoritative_remarks():
    file = prepare_file(workbook_bytes(), "客户.xlsx")
    source = page_source(storage_path(file["storage_key"]), 1)
    result = ExtractionResult(lines=[ExtractedLine(product="项目1", type_note="错误备注", difficulty="A",
        evidence=[Evidence(file_id=1, page=1, sheet="客户明细", cell="B2", text="项目1")])])
    result = validate_excel_rows(result, {(1, 1): source})
    assert len(result.lines) == 3
    assert {line.product for line in result.lines} == {"项目1", "项目2", "项目3"}
    assert all(line.type_note == "发光字" and line.difficulty == "B" for line in result.lines)
    assert all(line.difficulty_note == "B" for line in result.lines)
    assert result.lines[1].quantity == 2
    assert result.lines[1].uncertainties
    assert any(e.cell == "D2" and e.text == "B" for e in result.lines[0].evidence)
    assert len(validate_excel_rows(result, {(1, 1): source}).lines) == 3


def test_excel_source_and_image_endpoints_auth_and_provenance(client):
    response = client.post("/api/jobs", data={"title": "Excel测试"},
                           files={"files": ("客户.xlsx", workbook_bytes())})
    assert response.status_code == 201, response.text
    job = response.json()
    file_id = job["files"][0]["id"]
    source = client.get(f"/api/customer-files/{file_id}/source/1").json()
    assert "storage" not in source["images"][0]
    url = source["images"][0]["url"]
    assert client.get(url).status_code == 200
    for path in (url, f"/api/customer-files/{file_id}/source/1"):
        assert client.get(path, headers={"X-Admin-Token": "bad"}).status_code == 401
    assert client.get(f"/api/customer-files/{file_id}/images/original").status_code == 404
    line = {"id": "excel", "product": "项目1", "evidence": [
        {"file_id": file_id, "page": 1, "sheet": "客户明细", "cell": "B200", "text": "不存在"}]}
    update = {"revision": job["revision"], "reason": "核对来源", "lines": [line]}
    assert client.put(f'/api/jobs/{job["id"]}/requirements', json=update).status_code == 422
    line["evidence"][0]["cell"] = "B2"
    assert client.put(f'/api/jobs/{job["id"]}/requirements', json=update).status_code == 200


@pytest.mark.parametrize("part,transform", [
    ("xl/workbook.xml", lambda value: b"<!DOCTYPE foo [<!ENTITY x 'bad'>]>" + value),
    ("xl/styles.xml", lambda value: b"<!DOCTYPE foo [<!ENTITY x 'bad'>]>" + value),
    ("xl/worksheets/sheet1.xml", lambda value: value.replace(b'r="A2"', b'r="A99999"')),
])
def test_unsafe_xml_or_out_of_bounds_cells_are_rejected_and_cleaned(part, transform):
    with pytest.raises(FileRejected):
        prepare_file(modify_zip(workbook_bytes(drawing=False), part, transform), "bad.xlsx")
    assert not list(settings().storage_dir.iterdir())


def test_macros_rejected_but_excel_pages_are_not_capped(monkeypatch):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(workbook_bytes())) as original, zipfile.ZipFile(out, "w") as output:
        for part in original.infolist():
            output.writestr(part.filename, original.read(part.filename))
        output.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(FileRejected, match="宏"):
        prepare_file(out.getvalue(), "bad.xlsm")
    assert not list(settings().storage_dir.iterdir())
    monkeypatch.setattr(settings(), "max_job_pages", 1)
    file = prepare_file(workbook_bytes(), "many.xlsx")
    assert file["page_count"] == 2


def test_formula_cache_never_executes():
    def set_cache(value):
        tree = ET.fromstring(value)
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for cell in tree.findall(".//s:c", ns):
            if cell.find("s:f", ns) is not None:
                cell.find("s:v", ns).text = "999"
        return ET.tostring(tree)
    data = modify_zip(workbook_bytes(drawing=False), "xl/worksheets/sheet1.xml",
                      set_cache)
    file = prepare_file(data, "cache.xlsx")
    source = page_source(storage_path(file["storage_key"]), 1)
    formula = source["rows"][1]["cells"][-1]
    assert formula["value"] == "999"
    assert formula["formula_status"] == "cached_unverified"


def test_unstructured_difficulty_note_is_preserved_without_guessing():
    source = {"sheet": "明细", "header_rows": [{"row": 1, "cells": [
        {"cell": "A1", "value": "产品名称"}, {"cell": "B1", "value": "工艺难度备注"}]}],
        "rows": [{"row": 2, "cells": [{"cell": "A2", "value": "招牌"},
                                      {"cell": "B2", "value": "异形转角，须工厂核实"}]}], "images": []}
    line = validate_excel_rows(ExtractionResult(lines=[]), {(1, 1): source}).lines[0]
    assert line.difficulty == "unknown"
    assert line.difficulty_note == "异形转角，须工厂核实"


def test_model_excel_input_drawings_and_omission_fallback(monkeypatch):
    monkeypatch.setattr(settings(), "excel_native_first", False)
    file = SimpleNamespace(id=41, **prepare_file(workbook_bytes(), "customer.xlsx"))
    seen = []

    def handler(request):
        parts = json.loads(request.content)["messages"][1]["content"]
        page = len(seen) + 1
        seen.append(parts)
        assert any("原始文本/单元格" in part.get("text", "") for part in parts)
        assert len([part for part in parts if part["type"] == "image_url"]) == (1 if page == 1 else 0)
        result = {"lines": [{"source_item": "1", "product": "项目1" if page == 1 else "立牌",
            "difficulty": "A", "evidence": [{"file_id": 41, "page": page,
                "sheet": "客户明细" if page == 1 else "造型", "cell": "B2", "text": "项目"}]}]}
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop",
            "message": {"content": json.dumps(result)}}]})
    configure_transport(monkeypatch, handler)
    result, attempts, _ = deepseek.extract([file], "", [], "")
    assert attempts == 3 and len(result.lines) == 4
    assert result.lines[0].difficulty == "B"
    assert result.lines[-1].type_note == "双面"
    assert not any("跨页出现相同" in issue for line in result.lines for issue in line.uncertainties)


@pytest.mark.parametrize("count", [13, 61, 101])
def test_drawing_count_is_not_limited_and_all_originals_are_retained(count):
    file = prepare_file(workbook_bytes(image_count=count), "many-images.xlsx")
    directory = storage_path(file["storage_key"])
    source = page_source(directory, 1)
    assert len(source["images"]) == count
    assert len({image["storage"] for image in source["images"]}) == count
    assert all((directory / image["storage"]).is_file() for image in source["images"])
    assert not (directory / "page-1.png").exists()
    with Image.open(ensure_preview(directory, 1)) as preview:
        assert preview.width * preview.height < 25_000_000
    batches = list(image_batches(directory, 1))
    assert batches[0][0] == 0 and batches[-1][1] == count
    assert all(end - start <= 6 for start, end in batches)
    assert all(batches[index][1] == batches[index + 1][0] for index in range(len(batches) - 1))


def test_three_digit_image_names_are_available_and_authenticated(client):
    job = client.post("/api/jobs", data={"title": "101张图纸"},
                      files={"files": ("many.xlsx", workbook_bytes(image_count=101))}).json()
    file_id = job["files"][0]["id"]
    source = client.get(f"/api/customer-files/{file_id}/source/1").json()
    assert len(source["images"]) == 101
    last = source["images"][-1]["url"]
    assert last.endswith("excel-image-101.png")
    assert client.get(last).status_code == 200
    assert client.get(last, headers={"X-Admin-Token": "bad"}).status_code == 401
    assert client.get(last.replace("101.png", "102.png")).status_code == 404


def test_byte_budget_splits_even_few_large_drawings(monkeypatch):
    from app import customer_excel
    file = prepare_file(workbook_bytes(image_count=3), "large-drawings.xlsx")
    monkeypatch.setattr(customer_excel, "IMAGE_BATCH_BYTES", 1)
    assert list(image_batches(storage_path(file["storage_key"]), 1)) == [(0, 1), (1, 2), (2, 3)]


def test_more_than_twenty_excel_pages_can_upload_append_extract_and_keep_provenance(client, monkeypatch):
    monkeypatch.setattr(settings(), "max_job_pages", 1)
    response = client.post("/api/jobs", data={"title": "长Excel"},
        files={"files": ("long.xlsx", workbook_bytes(rows=400, drawing=False))})
    assert response.status_code == 201, response.text
    job = response.json()
    file_id = job["files"][0]["id"]
    assert job["files"][0]["page_count"] == 22
    assert client.get(f"/api/customer-files/{file_id}/source/21").status_code == 200
    response = client.put(f'/api/jobs/{job["id"]}/requirements', json={
        "revision": job["revision"], "reason": "核对末行来源", "lines": [{
            "id": "last", "product": "项目400",
            "evidence": [{"file_id": file_id, "page": 21, "sheet": "客户明细", "cell": "B401", "text": "项目400"}]}]})
    assert response.status_code == 200, response.text
    job = response.json()
    response = client.post(f'/api/jobs/{job["id"]}/files', data={"revision": job["revision"]},
        files={"files": ("extra.xlsx", workbook_bytes(rows=2, drawing=False))})
    assert response.status_code == 200, response.text
    job = response.json()
    monkeypatch.setattr(deepseek, "get_key", lambda: "fake-key")
    monkeypatch.setattr(deepseek, "extract", lambda *args: (ExtractionResult(lines=[]), 1, {}))
    response = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": job["revision"], "allow_external_processing": True})
    assert response.status_code == 200, response.text


def test_all_drawing_batches_are_read_without_duplicate_projects_or_lost_conflicts(monkeypatch):
    monkeypatch.setattr(settings(), "excel_native_first", False)
    file = SimpleNamespace(id=41, **prepare_file(workbook_bytes(image_count=13), "many.xlsx"))
    seen, image_names = [], []

    def handler(request):
        parts = json.loads(request.content)["messages"][1]["content"]
        source = next(json.loads(part["text"].split(":\n", 1)[1]) for part in parts
                      if "原始文本/单元格" in part.get("text", ""))
        images = [part for part in parts if part["type"] == "image_url"]
        page = 1 if source["sheet"] == "客户明细" else 2
        seen.append((page, len(images)))
        image_names.extend(image["storage"] for image in source["images"])
        width = 320 if len(seen) == 2 else 300
        line = {"source_item": "1", "product": "项目1" if page == 1 else "立牌", "quantity": 2,
            "evidence": [{"file_id": 41, "page": page, "sheet": source["sheet"], "cell": "B2", "text": "项目"}]}
        if page == 1:
            line["dimension_evidence"] = [{"file_id": 41, "page": 1, "sheet": source["sheet"], "cell": "G2",
                "field": "width_mm", "value": width, "source": "image", "text": f"{width}mm"}]
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop",
            "message": {"content": json.dumps({"lines": [line]})}}], "usage": {"total_tokens": 10}})
    configure_transport(monkeypatch, handler)
    result, attempts, usage = deepseek.extract([file], "", [], "")
    assert seen == [(1, 6), (1, 6), (1, 1), (2, 0)]
    assert image_names == [f"excel-image-{index}.png" for index in range(1, 14)]
    assert attempts == 5 and usage["total_tokens"] == 40
    assert len(result.lines) == 4
    assert result.lines[0].quantity == 2 and result.lines[0].width_mm is None
    assert result.lines[0].uncertainties
    assert {line.product for line in result.lines} == {"项目1", "项目2", "项目3", "立牌"}


def test_conflicting_batch_quantities_stay_unknown_even_when_later_views_repeat():
    results = [ExtractionResult(lines=[ExtractedLine(product="牌", source_item="1", quantity=count,
        evidence=[Evidence(file_id=1, page=1, sheet="明细", cell="B2", text=f"{count}件")])])
        for count in (2, 3, 2)]
    result = merge_drawing_results(results)
    assert len(result.lines) == 1 and result.lines[0].quantity is None
    assert any("quantity不一致" in item for item in result.lines[0].uncertainties)


def test_failed_drawing_batch_keeps_paid_usage_and_stops_later_batches(monkeypatch):
    monkeypatch.setattr(settings(), "excel_native_first", False)
    file = SimpleNamespace(id=41, **prepare_file(workbook_bytes(image_count=13), "many.xlsx"))
    calls = []

    def handler(request):
        calls.append(request)
        if len(calls) == 2:
            return httpx.Response(401, json={"error": "bad key"})
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop",
            "message": {"content": '{"lines":[]}'}}], "usage": {"total_tokens": 10}})
    configure_transport(monkeypatch, handler)
    with pytest.raises(deepseek.ModelFailure) as failure:
        deepseek.extract([file], "", [], "")
    assert len(calls) == 2
    assert failure.value.attempts == 3 and failure.value.usage["total_tokens"] == 10
    assert failure.value.diagnostics["image_start"] == 7
    assert failure.value.diagnostics["image_end"] == 12
