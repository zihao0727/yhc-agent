import copy
import io
import json
import threading
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from PIL import Image
from sqlalchemy import select

from app import agent_worker, deepseek
from app.config import settings
from app.customer_files import prepare_file
from app.extraction_checkpoints import ExtractionCheckpoint, ExtractionYield
from app.models import AgentRun
from app.requirements_schemas import ExtractedLine
from app.source_agent import SourceSession, covers_image
from app.source_agent_tools import SPECS
from app.source_reconciliation import source_review
from test_agent_loop import call, provider, start
from test_large_pdf import document_bytes
from test_redis_agent import redis_worker  # noqa: F401
from test_requirements import create_job, line_input, png_bytes, save_lines


@pytest.fixture(autouse=True)
def enable_source_agent(isolate_customer_storage, monkeypatch):
    monkeypatch.setattr(settings(), "source_agent_enabled", True)


def file_fixture(file_id=1, data=None, filename="original.png"):
    return SimpleNamespace(id=file_id, **prepare_file(data or png_bytes(), filename))


def execute(session, name, **args):
    return session.execute(name, SPECS[name][0].model_validate(args))


def line(file_id=1, product="Sign", **fields):
    return {"product": product, "quantity": 2, "evidence": [
        {"file_id": file_id, "page": 1, "text": product, "source": "image"}], **fields}


def save(session, raw, identifier="", reason="Read original"):
    return execute(session, "save_observations", observations=[
        {"line": raw, "observation_id": identifier, "reason": reason}])


def review_calls(file_ids):
    return [call("begin_source_review", summary="Compare originals for omissions and duplicate views"),
            *[call("read_source", file_id=fid, page=1, purpose="review") for fid in file_ids],
            call("record_source_reviews", pages=[{
                "file_id": fid, "page": 1, "role": "drawing", "status": "reviewed", "basis": "reread",
                "reason": "Compared original and saved observations"} for fid in file_ids]),
            call("finish_source_review", summary="Ready for pricing")]


def coverage_calls(file_ids):
    return [call("record_source_reviews", pages=[{
                "file_id": fid, "page": 1, "role": "drawing", "status": "reviewed",
                "reason": "Checked page coverage during initial reading"} for fid in file_ids]),
            call("finish_source_review", summary="Ready for pricing")]


def checkpoint(saved=None):
    snapshots, charges = [], []
    cache = ExtractionCheckpoint("ai", saved or {},
        lambda state, usage: (snapshots.append(copy.deepcopy(state)), charges.append(copy.deepcopy(usage))),
        lambda: None)
    return cache, snapshots, charges


def test_default_outer_route_uses_tool_agent_and_direct_pixels(monkeypatch):
    seen = provider(monkeypatch, [
        call("set_reading_plan", plan="Read the schedule and check actual products"),
        call("read_source", file_id=1, page=1),
        call("save_observations", observations=[{"line": line(), "reason": "Visible product"}]),
        *coverage_calls([1])])
    result, attempts, usage = deepseek.extract([file_fixture()], "", [], "")
    assert attempts == len(seen) == 5 and usage["total_tokens"] == 50
    assert len(result.lines) == 1 and result.lines[0].measurement_version == 3
    assert result.source_agent["read_pages"] == result.source_agent["reviewed_pages"] == 1
    assert not source_review(result.model_dump(mode="json"))["required"]
    assert any(part.get("type") == "image_url" for message in seen[2]
               if isinstance(message.get("content"), list) for part in message["content"])


def test_ai_can_correct_bad_code_and_exclude_catalog_without_preserving_false_rows():
    file = file_fixture(filename="directory.pdf", data=document_bytes(text=True))
    session = SourceSession([file], "", [], "", None)
    execute(session, "read_source", file_id=1, page=1)
    save(session, line(product="Cover"))
    # The code is an AI observation; native text need not contain its exact spelling.
    save(session, line(product="4F plan", source_code="4F-BD-02"))
    save(session, line(product="Actual sign", source_code="SN-02"), "o000002", "Corrected after looking at image")
    execute(session, "reconcile_observations", action="exclude", observation_ids=["o000001"],
            reason="Cover is a directory entry, not a purchased product")
    result = session.result()
    assert [x.product for x in result.lines] == ["Actual sign"]
    assert result.source_agent["archived"]["o000001"]["product"] == "Cover"
    assert [d["action"] for d in result.source_agent["decisions"]] == ["add", "add", "revise", "exclude"]
    execute(session, "reconcile_observations", action="restore", observation_ids=["o000001"],
            reason="Reconsider original evidence")
    assert len(session.result().lines) == 2


def test_ai_can_merge_cross_file_different_names_without_code_equality_or_quantity_sum():
    session = SourceSession([file_fixture(), file_fixture(2)], "", [], "", None)
    for fid in (1, 2):
        execute(session, "read_source", file_id=fid, page=1)
    save(session, line(product="Entrance identification", quantity=6))
    save(session, line(2, "Front elevation", quantity=1))
    merged = line(product="Entrance identification", quantity=6, width_mm=300, height_mm=120)
    merged["evidence"].extend(line(2, "Front elevation")["evidence"])
    execute(session, "reconcile_observations", action="merge", observation_ids=["o000001", "o000002"],
            merged_line=merged, reason="The drawing title and location identify the schedule product")
    result = session.result()
    assert len(result.lines) == 1 and result.lines[0].quantity == 6
    assert {e.file_id for e in result.lines[0].evidence} == {1, 2}
    assert result.source_links[0].status == "merged"
    assert session.state["decisions"][-1]["before"]["o000002"]["quantity"] == 1


def test_observations_cannot_invent_order_files_or_claim_unseen_images():
    session = SourceSession([file_fixture()], "", [], "", None)
    with pytest.raises(ValueError, match="先读取"):
        save(session, line())
    with pytest.raises(ValueError, match="不属于"):
        execute(session, "read_source", file_id=999, page=1)
    # Inspecting an image is an actual tool action, not an AI-supplied verification flag.
    execute(session, "read_source", file_id=1, page=1)
    save(session, line(quantity_evidence=[{"file_id": 1, "page": 1, "text": "2", "quantity": 2,
        "validation_status": "verified", "validation_note": "passed"}]))
    assert session.result().lines[0].quantity_evidence[0].validation_status == "unchecked"


def test_review_requires_real_reread_and_changes_invalidate_prior_conclusion():
    session = SourceSession([file_fixture()], "", [], "", None)
    execute(session, "read_source", file_id=1, page=1)
    save(session, line())
    execute(session, "begin_source_review", summary="Check omissions")
    record = {"file_id": 1, "page": 1, "role": "drawing", "status": "reviewed",
              "basis": "reread", "reason": "Compared original"}
    with pytest.raises(ValueError, match="purpose=review"):
        execute(session, "record_source_reviews", pages=[record])
    execute(session, "read_source", file_id=1, page=1, purpose="review")
    execute(session, "record_source_reviews", pages=[record])
    assert session.audit()["reviewed_pages"] == 1
    save(session, line(quantity=3), "o000001", "Correct quantity after rereading")
    assert session.audit()["reviewed_pages"] == 0
    execute(session, "record_source_reviews", pages=[record])
    assert session.audit()["reviewed_pages"] == 1


def test_initial_page_conclusions_cover_many_items_without_per_item_checks_or_second_read(monkeypatch):
    observations = [{"line": line(product=f"Sign {i}"), "reason": "Visible ordered product"} for i in range(6)]
    seen = provider(monkeypatch, [
        call("read_source", file_id=1, page=1),
        call("save_observations", observations=observations[:3]),
        call("save_observations", observations=observations[3:]),
        *coverage_calls([1]),
    ])
    result, attempts, usage = deepseek.extract([file_fixture()], "", [], "")
    assert attempts == len(seen) == 5 and usage["total_tokens"] == 50
    assert len(result.lines) == 6 and result.source_agent["reviewed_pages"] == 1
    assert len([read for read in result.source_agent["reads"] if read["medium"] == "image"]) == 1
    assert result.source_agent["page_reviews"][0]["basis"] == "initial_read"
    assert not source_review(result.model_dump(mode="json"))["required"]
    # The original remains available for subsequent saves on a dense page.
    assert any(isinstance(message.get("content"), list) for message in seen[2])
    assert not any(isinstance(message.get("content"), list) for message in seen[4])
    assert not json.loads(seen[2][1]["content"])["pages_awaiting_saved_conclusion"]


def test_selective_review_reuses_initial_read_and_only_reopens_questionable_region():
    file = file_fixture()
    session = SourceSession([file], "", [], "", None)
    execute(session, "read_source", file_id=1, page=1)
    save(session, line(quantity=None, uncertainties=["Quantity label is unclear"]))
    record = {"file_id": 1, "page": 1, "role": "drawing", "status": "unresolved",
              "reason": "Top-right quantity label needs enlargement"}
    execute(session, "record_source_reviews", pages=[record])
    assert session.audit()["reviewed_pages"] == 0
    assert json.loads(session.messages()[1]["content"])["source_questions"][0]["reason"] == record["reason"]
    execute(session, "begin_source_review", summary="Only inspect unclear quantity")
    assert session.state["reviews"]["1:1"]["reason"] == record["reason"]
    execute(session, "inspect_source_region", file_id=1, page=1, region=[.5, 0, 1, .5], purpose="review")
    save(session, line(quantity=3), "o000001", "Quantity clear in enlarged region")
    execute(session, "record_source_reviews", pages=[{
        **record, "status": "reviewed", "reason": "Initial page coverage plus enlarged quantity label"}])
    execute(session, "finish_source_review", summary="Only uncertain region reread")
    assert session.audit()["reviewed_pages"] == 1
    assert not session.delivered(file, 1, review=True)
    assert len([read for read in session.state["reads"].values() if read["medium"] == "image"]) == 2


def test_initial_conclusions_do_not_claim_unread_pages_and_explicit_reread_remains_truthful():
    session = SourceSession([file_fixture(), file_fixture(2)], "", [], "", None)
    record = {"file_id": 1, "page": 1, "role": "drawing", "status": "reviewed", "reason": "Page checked"}
    with pytest.raises(ValueError, match="先读取"):
        execute(session, "record_source_reviews", pages=[record])
    execute(session, "read_source", file_id=1, page=1)
    execute(session, "record_source_reviews", pages=[record])
    execute(session, "begin_source_review", summary="Check remaining issues")
    assert session.audit()["reviewed_pages"] == 1
    with pytest.raises(ValueError, match="purpose=review"):
        execute(session, "record_source_reviews", pages=[{**record, "basis": "reread"}])
    execute(session, "finish_source_review", summary="Second file still unread")
    assert session.audit()["unread_pages"] == [{"file_id": 2, "page": 1}]
    assert session.result().source_coverage[1].status == "unresolved"


def test_legacy_review_conclusions_still_require_actual_reread():
    file = file_fixture()
    cache, _, _ = checkpoint()
    session = SourceSession([file], "", [], "", cache)
    execute(session, "read_source", file_id=1, page=1)
    session.state["reviews"]["1:1"] = {
        "file_id": 1, "page": 1, "role": "drawing", "status": "reviewed", "reason": "Legacy conclusion"}
    cache.source_agent = copy.deepcopy(session.state)
    resumed = SourceSession([file], "", [], "", cache)
    assert resumed.audit()["reviewed_pages"] == 0
    execute(resumed, "begin_source_review", summary="Recheck original")
    execute(resumed, "read_source", file_id=1, page=1, purpose="review")
    assert resumed.audit()["reviewed_pages"] == 1


def test_initial_conclusion_changes_invalidate_coverage_and_resume_keeps_other_pages():
    files = [file_fixture(), file_fixture(2)]
    cache, _, _ = checkpoint()
    session = SourceSession(files, "", [], "", cache)
    for fid in (1, 2):
        execute(session, "read_source", file_id=fid, page=1)
        save(session, line(fid))
    records = [{"file_id": fid, "page": 1, "role": "drawing", "status": "reviewed",
                "reason": "Initial page coverage checked"} for fid in (1, 2)]
    execute(session, "record_source_reviews", pages=records)
    cache.source_agent = copy.deepcopy(session.state)
    resumed = SourceSession(files, "", [], "", cache)
    save(resumed, line(quantity=3), "o000001", "Quantity corrected from another file")
    execute(resumed, "begin_source_review", summary="Only changed page needs a new conclusion")
    assert resumed.audit()["reviewed_pages"] == 1
    assert resumed.audit()["unreviewed_pages"] == [{"file_id": 1, "page": 1}]
    execute(resumed, "record_source_reviews", pages=records[:1])
    assert resumed.audit()["reviewed_pages"] == 2
    assert len(resumed.state["reads"]) == len(session.state["reads"])


def test_initial_pdf_coverage_still_requires_original_pixels():
    session = SourceSession([file_fixture(filename="original.pdf", data=document_bytes(text=True))],
                            "", [], "", None)
    execute(session, "read_source", file_id=1, page=1, include_image=False)
    execute(session, "record_source_reviews", pages=[{
        "file_id": 1, "page": 1, "role": "drawing", "status": "reviewed", "reason": "Only text seen"}])
    assert session.audit()["reviewed_pages"] == 0
    assert session.audit()["unread_pages"] == [{"file_id": 1, "page": 1}]
    execute(session, "inspect_source_region", file_id=1, page=1)
    assert session.audit()["reviewed_pages"] == 1


def test_observation_summaries_paginate_without_replaying_full_evidence():
    session = SourceSession([file_fixture()], "", [], "", None)
    execute(session, "read_source", file_id=1, page=1)
    for i in range(20):
        save(session, line(product=f"Sign {i}", notes="Original process detail " * 80))
    summary = execute(session, "list_observations")
    assert len(summary["observations"]) == 15 and summary["next_offset"] == 15
    assert summary["observations"][0]["line"]["quantity"] == 2
    assert summary["observations"][0]["line"]["sources"] == [{"file_id": 1, "page": 1}]
    assert "notes" not in summary["observations"][0]["line"]
    detail = execute(session, "list_observations", observation_ids=["o000001"], view="full")
    assert len(detail["observations"]) == 1 and detail["observations"][0]["line"]["notes"]
    assert len(json.dumps(detail["observations"][0])) > len(json.dumps(summary["observations"][0])) * 5
    assert execute(session, "list_sources")["pages"] == []
    assert len(execute(session, "list_sources", status="unreviewed")["pages"]) == 1
    assert len(execute(session, "list_sources", status="all")["pages"]) == 1


def test_context_guides_repeated_queries_to_bulk_page_conclusions():
    session = SourceSession([file_fixture()], "", [], "", None)
    execute(session, "read_source", file_id=1, page=1)
    save(session, line())
    session.state["history"] = [{"tool": tool} for tool in (
        "list_sources", "list_observations", "list_source_notes")]
    message = json.loads(session.messages()[1]["content"])
    assert not message["pages_awaiting_saved_conclusion"]
    assert message["pages_awaiting_coverage_conclusion"][0]["saved_items"] == 1
    assert "record_source_reviews" in message["next_action"]
    assert message["repeated_queries"]["count"] == 3


def test_region_pixels_do_not_claim_unseen_half_is_reviewed():
    session = SourceSession([file_fixture()], "", [], "", None)
    execute(session, "begin_source_review", summary="Check both halves")
    execute(session, "inspect_source_region", file_id=1, page=1, region=[0, 0, 1, .5], purpose="review")
    execute(session, "record_source_reviews", pages=[{
        "file_id": 1, "page": 1, "role": "drawing", "status": "reviewed", "reason": "Only top half checked"}])
    assert session.audit()["reviewed_pages"] == 0
    assert session.result().source_coverage[0].status == "unresolved"
    execute(session, "inspect_source_region", file_id=1, page=1, region=[0, .5, 1, 1], purpose="review")
    assert session.audit()["reviewed_pages"] == 1
    assert covers_image([[0, 0, 1, .5], [0, .5, 1, 1]])
    assert not covers_image([[0, 0, .49, 1], [.5, 0, 1, 1]])


def test_pdf_text_truncation_and_search_never_count_as_complete_read():
    file = file_fixture(filename="drawing.pdf", data=document_bytes(pages=2, text=True))
    session = SourceSession([file], "", [], "", None)
    session.texts[1, 1] = session.text(1, 1) * 5
    execute(session, "search_sources", file_id=1, query="TOP")
    assert session.audit()["read_pages"] == 0
    execute(session, "begin_source_review", summary="Check original")
    page = execute(session, "read_source", file_id=1, page=1, limit=100, purpose="review")
    assert page["next_offset"] is not None
    execute(session, "record_source_reviews", pages=[{
        "file_id": 1, "page": 1, "role": "drawing", "status": "reviewed", "reason": "Partial text seen"}])
    assert session.audit()["reviewed_pages"] == 0
    execute(session, "read_source", file_id=1, page=1, offset=page["next_offset"], purpose="review")
    assert session.audit()["reviewed_pages"] == 1
    execute(session, "finish_source_review", summary="Some pages still unread")
    audit = source_review(session.result().model_dump(mode="json"))
    assert audit["required"] and audit["unread_pages"] == 1


def test_excel_reads_actual_cells_and_embedded_image_crops():
    from openpyxl.drawing.image import Image as ExcelImage
    workbook = Workbook()
    workbook.active.append(["Name", "Quantity"])
    workbook.active.append(["Sign", 2])
    workbook.active.add_image(ExcelImage(io.BytesIO(png_bytes())), "C2")
    data = io.BytesIO()
    workbook.save(data)
    file = file_fixture(data=data.getvalue(), filename="schedule.xlsx")
    session = SourceSession([file], "", [], "", None)
    execute(session, "begin_source_review", summary="Check cells and embedded drawing")
    result = execute(session, "read_source", file_id=1, page=1, purpose="review")
    assert result["images"][0]["image_index"] == 0
    assert not session.delivered(file, 1, review=True)
    execute(session, "inspect_source_region", file_id=1, page=1, image_index=0,
            region=[0, 0, .5, 1], purpose="review")
    assert not session.delivered(file, 1, review=True)
    execute(session, "inspect_source_region", file_id=1, page=1, image_index=0, purpose="review")
    assert session.delivered(file, 1, review=True)
    sheet = result["sheet"]
    good = line(evidence=[{"file_id": 1, "page": 1, "sheet": sheet, "cell": "A2", "text": "Sign"}])
    save(session, good)
    with pytest.raises(ValueError, match="单元格不存在"):
        save(session, line(evidence=[{"file_id": 1, "page": 1, "sheet": sheet, "cell": "X999", "text": "Sign"}]))


def test_excel_image_selection_and_absent_images_do_not_claim_a_view():
    from openpyxl.drawing.image import Image as ExcelImage
    workbook = Workbook()
    workbook.active.append(["Name", "Quantity"])
    workbook.active.append(["Sign", 2])
    data = io.BytesIO()
    workbook.save(data)
    file = file_fixture(data=data.getvalue(), filename="schedule.xlsx")
    session = SourceSession([file], "", [], "", None)
    result = execute(session, "inspect_source_region", file_id=1, page=1)
    assert result["status"] == "no_embedded_images"
    assert result["available_images"] == []
    assert not result["viewed"] and not result["images_follow"]
    assert session.audit()["read_pages"] == 0

    workbook.active.add_image(ExcelImage(io.BytesIO(png_bytes())), "C2")
    data = io.BytesIO()
    workbook.save(data)
    file = file_fixture(file_id=2, data=data.getvalue(), filename="illustrated.xlsx")
    session = SourceSession([file], "", [], "", None)
    result = execute(session, "inspect_source_region", file_id=2, page=1)
    assert result["status"] == "image_selection_required"
    assert result["available_images"][0]["image_index"] == 0
    assert result["available_images"][0]["cell"] == "C2"
    assert not result["viewed"] and session.audit()["read_pages"] == 0


def test_pricing_does_not_queue_absent_excel_images_for_replay(test_db):
    from app.models import RequirementJob, CustomerFile
    from app.requirements_schemas import RequirementLine
    from app.source_agent_tools import InspectRegion
    from app.agent_execution import execute_tool
    from app import agent_work_state as work
    workbook = Workbook()
    workbook.active.append(["Name", "Quantity"])
    workbook.active.append(["Sign", 2])
    output = io.BytesIO()
    workbook.save(output)
    with test_db.begin() as db:
        job = RequirementJob(title="No images", requirements=[
            RequirementLine(id="one", product="Sign").model_dump(mode="json")])
        db.add(job)
        db.flush()
        file = CustomerFile(job_id=job.id, **prepare_file(output.getvalue(), "cells.xlsx"))
        db.add(file)
        db.flush()
        run = AgentRun(job_id=job.id, model="test", steps=[],
                       usage={"durable": True, "active_line_id": "one", "pricing_concurrency": 4})
        db.add(run)
        args = InspectRegion(file_id=file.id, page=1)
        work.update_state(job, run, "one", source_view_pending=[args.model_dump(mode="json")])
        result = execute_tool(db, job, run, "view_source_region", args, set())
        assert result["status"] == "no_embedded_images"
        assert "read_source" in result["next_action"]
        assert not result.get("_source_images")
        assert work.line_state(job, run, "one")["source_view_pending"] == []


def test_same_excel_row_cannot_be_added_again_and_missing_rows_survive_review_flag():
    workbook = Workbook()
    workbook.active.append(["序号", "产品名称", "数量", "单位"])
    workbook.active.append([1, "Sign", 2, "套"])
    workbook.active.append([2, "Other sign", 3, "套"])
    data = io.BytesIO()
    workbook.save(data)
    file = file_fixture(data=data.getvalue(), filename="schedule.xlsx")
    session = SourceSession([file], "", [], "", None)
    page = execute(session, "read_source", file_id=1, page=1)
    raw = line(document_role="requirements", evidence=[
        {"file_id": 1, "page": 1, "sheet": page["sheet"], "cell": "B2", "text": "Sign"}])
    save(session, raw)
    with pytest.raises(ValueError, match="o000001"):
        save(session, raw)
    assert len(session.state["observations"]) == 1
    execute(session, "record_source_reviews", pages=[{
        "file_id": 1, "page": 1, "role": "requirements", "status": "reviewed",
        "reason": "Claimed all rows complete"}])
    finished = execute(session, "finish_source_review", summary="Done")
    assert not finished["finished"] and session.state["phase"] == "review"
    result = session.result()
    audit = source_review(result.model_dump(mode="json"))
    assert audit["required"] and audit["inventory_issues"][0]["row"] == 3
    save(session, {**raw, "product": "Other sign", "evidence": [
        {"file_id": 1, "page": 1, "sheet": page["sheet"], "cell": "B3", "text": "Other sign"}]})
    assert not session.inventory()["issues"]
    save(session, {**raw, "width_mm": 100}, "o000001")
    assert len(session.state["observations"]) == 2
    assert execute(session, "finish_source_review", summary="Rows reconciled")["finished"]


def test_unmatched_drawings_remain_visible_with_explicit_scope_question():
    session = SourceSession([file_fixture(), file_fixture(2)], "", [], "", None)
    for fid in (1, 2):
        execute(session, "read_source", file_id=fid, page=1)
    save(session, line(2, "Drawing detail", document_role="drawing"))
    save(session, line(1, "Ordered sign", document_role="requirements"))
    result = session.result()
    assert [item.product for item in result.lines] == ["Drawing detail", "Ordered sign"]
    assert result.source_agent["observation_ids"] == ["o000001", "o000002"]
    assert result.source_links[0].status == "unresolved"
    assert result.source_links[0].product == "Drawing detail"
    assert result.source_agent["unresolved_drawings"][0]["observation_id"] == "o000001"
    assert result.questions
    save(session, line(2, "Drawing detail", document_role="requirements"), "o000001")
    assert len(session.result().lines) == 2
    assert not session.result().source_links
    assert not session.audit()["unresolved_drawings"]


def test_completed_legacy_checkpoint_reopens_to_recover_missing_original_row(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["序号", "产品名称", "数量", "单位"])
    workbook.active.append([1, "Sign", 2, "套"])
    workbook.active.append([2, "Other sign", 3, "套"])
    data = io.BytesIO()
    workbook.save(data)
    file = file_fixture(data=data.getvalue(), filename="schedule.xlsx")
    session = SourceSession([file], "", [], "", None)
    page = execute(session, "read_source", file_id=1, page=1)
    save(session, line(document_role="requirements", evidence=[
        {"file_id": 1, "page": 1, "sheet": page["sheet"], "cell": "B2", "text": "Sign"}]))
    session.state["phase"] = "complete"
    cache, _, _ = checkpoint({"signature": "ai", "source_agent": session.state})
    seen = provider(monkeypatch, [
        call("save_observations", observations=[{"line": line(product="Other sign", quantity=3,
            document_role="requirements", evidence=[{
                "file_id": 1, "page": 1, "sheet": page["sheet"], "cell": "B3", "text": "Other sign"}]),
            "reason": "Recovered missing original row"}]),
        call("record_source_reviews", pages=[{
            "file_id": 1, "page": 1, "role": "requirements", "status": "reviewed",
            "reason": "Both original rows accounted for"}]),
        call("finish_source_review", summary="Complete original inventory")])
    result, attempts, _ = deepseek.extract([file], "", [], "", _checkpoint=cache)
    assert attempts == len(seen) == 3
    assert len(result.lines) == 2 and not result.source_agent["inventory"]["issues"]
    assert not source_review(result.model_dump(mode="json"))["required"]


def test_local_bad_read_does_not_block_other_sources(monkeypatch):
    seen = provider(monkeypatch, [
        call("read_source", file_id=999, page=1),
        call("read_source", file_id=1, page=1),
        call("save_observations", observations=[{"line": line(), "reason": "Visible sign"}]),
        *review_calls([1])])
    result, _, _ = deepseek.extract([file_fixture()], "", [], "")
    assert len(result.lines) == 1
    assert "不属于" in seen[1][-1]["content"]


def test_slice_resume_replays_pixels_and_preserves_observation_ids_and_usage(monkeypatch):
    monkeypatch.setattr(settings(), "agent_slice_rounds", 1)
    seen = provider(monkeypatch, [
        call("read_source", file_id=1, page=1),
        call("save_observations", observations=[{"line": line(), "reason": "Visible sign"}]),
        *review_calls([1])])
    file, saved, charges, result = file_fixture(), {}, [], None
    for _ in range(10):
        cache = ExtractionCheckpoint("resume", saved,
            lambda state, usage: (saved.update(copy.deepcopy(state)), charges.append(usage)), lambda: None)
        try:
            result, _, _ = deepseek.extract([file], "", [], "", _checkpoint=cache)
        except ExtractionYield:
            continue
        break
    assert result is not None and len(result.lines) == 1
    assert result.source_agent["observation_ids"] == ["o000001"]
    assert result.source_agent["reviewed_pages"] == 1
    assert sum(t.get("total_tokens", 0) for t in charges) == len(seen) * 10
    assert any(message.get("role") == "user" and isinstance(message.get("content"), list)
               and any(part.get("type") == "image_url" for part in message["content"])
               for message in seen[1])
    _, attempts, usage = deepseek.extract([file], "", [], "", _checkpoint=cache)
    assert attempts == 0 and usage == {}


def test_budget_pause_keeps_reading_records_and_is_not_silent_completion(monkeypatch):
    monkeypatch.setattr(settings(), "source_agent_max_tokens", 10)
    provider(monkeypatch, [call("read_source", file_id=1, page=1)])
    cache, snapshots, charges = checkpoint()
    with pytest.raises(deepseek.ModelFailure) as failure:
        deepseek.extract([file_fixture()], "", [], "", _checkpoint=cache)
    assert failure.value.diagnostics["stage"] == "source_agent_budget"
    assert snapshots[-1]["source_agent"]["phase"] == "reading"
    assert snapshots[-1]["source_agent"]["pending_sources"]
    assert sum(t.get("total_tokens", 0) for t in charges) == 10


def test_invalid_response_paid_usage_is_saved_for_retry(monkeypatch):
    def failed(*args):
        raise deepseek.AgentResponseFailure("truncated", {"finish_reason": "length"}, {"total_tokens": 17})
    monkeypatch.setattr(deepseek, "agent_turn", failed)
    cache, snapshots, charges = checkpoint()
    with pytest.raises(deepseek.ModelFailure):
        deepseek.extract([file_fixture()], "", [], "", _checkpoint=cache)
    assert snapshots[-1]["source_agent"]["total_tokens"] == 51
    assert snapshots[-1]["source_agent"]["response_recovery"]["active"]
    assert sum(t.get("total_tokens", 0) for t in charges) == 51


def test_human_confirmed_price_and_identity_cannot_be_replaced():
    raw = line_input(manual_unit_price="99.00", manual_price_note="Approved scope")
    raw["evidence"] = line()["evidence"]
    session = SourceSession([file_fixture()], "", [raw], "", None)
    identifier = next(iter(session.state["protected"]))
    execute(session, "read_source", file_id=1, page=1)
    with pytest.raises(ValueError, match="人工确认"):
        save(session, line(quantity=99), identifier)
    with pytest.raises(ValueError, match="人工确认"):
        execute(session, "reconcile_observations", action="exclude",
                observation_ids=[identifier], reason="AI changed its mind")
    assert session.result().source_agent["protected_lines"][identifier] == raw


def test_background_publication_can_remove_rows_and_resume_review_without_new_paid_reads(
        client, monkeypatch, redis_worker):
    job = create_job(client)
    fid = job["files"][0]["id"]
    seen = provider(monkeypatch, [
        call("read_source", file_id=fid, page=1),
        call("save_observations", observations=[
            {"line": line(fid, "Cover"), "reason": "Initial observation"},
            {"line": line(fid, "Sign"), "reason": "Initial observation"}]),
        call("reconcile_observations", action="exclude", observation_ids=["o000001"],
             reason="Reconsidered: cover is not an ordered product"),
        *review_calls([fid])])
    monkeypatch.setattr(settings(), "agent_background", True)
    monkeypatch.setattr(settings(), "agent_slice_rounds", 1)
    started = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": job["revision"], "allow_external_processing": True})
    assert started.status_code == 200, started.text
    snapshots = []
    for _ in range(10):
        agent_worker.run_pending_once(threading.Event())
        current = client.get(f'/api/jobs/{job["id"]}').json()
        snapshots.append(current)
        if current["status"] != "processing":
            break
    assert current["agent_runs"][0]["status"] == "succeeded", current["agent_runs"][0]["message"]
    assert not current["extraction"].get("incomplete")
    assert [r["product"] for r in current["requirements"]] == ["Sign"]
    two = next(s for s in snapshots if len(s["requirements"]) == 2)
    assert two["extraction"]["incomplete"]
    saved_id = two["requirements"][1]["id"]
    assert current["requirements"][0]["id"] == saved_id
    assert current["agent_runs"][0]["usage"]["total_tokens"] == len(seen) * 10
    report = client.get(f'/api/jobs/{job["id"]}/pipeline-audit').json()
    assert report["source_agent"]["decisions"][-1]["action"] == "exclude"
    assert report["source_agent"]["reviewed_pages"] == 1


def test_unread_source_keeps_priced_lines_as_partial_quote(client, test_db):
    from app.quote_engine import build_quote
    job = create_job(client)
    raw = line_input(manual_unit_price="100.00", manual_price_note="Manual quote")
    with test_db() as db:
        payload = build_quote(db, [raw], automatic=True, extraction={
            "source_agent": {"mode": "ai", "unread_pages": [{"file_id": job["files"][0]["id"], "page": 1}],
                             "unreviewed_pages": [{"file_id": job["files"][0]["id"], "page": 1}]}})
    assert payload["known_subtotal"] == "200.00"
    assert payload["total"] is None and not payload["complete"]
    assert payload["source_review"]["required"]


def test_mixed_attachments_are_reconciled_by_ai_then_produce_reviewable_quote(
        client, monkeypatch, price_input):
    from test_agent_loop import seed
    price = seed(client, price_input)
    workbook = Workbook()
    workbook.active.append(["Product", "Quantity"])
    workbook.active.append(["Entrance", 2])
    data = io.BytesIO()
    workbook.save(data)
    response = client.post("/api/jobs", data={"title": "Mixed AI quote", "business_rules_version": "legacy"},
        files=[("files", ("schedule.xlsx", data.getvalue())),
               ("files", ("drawing.pdf", document_bytes(text=True))),
               ("files", ("notes.png", png_bytes()))])
    assert response.status_code == 201, response.text
    job = response.json()
    excel_id, pdf_id, image_id = [f["id"] for f in job["files"]]
    schedule = line(excel_id, "Entrance", evidence=[{
        "file_id": excel_id, "page": 1, "sheet": "Sheet", "cell": "A2", "text": "Entrance"}])
    drawing = line(pdf_id, "Elevation")
    merged = {**schedule, "product": price["product"], "pricing_category": price["product"],
              "category_basis": "AI compared schedule and drawing",
              "material": price["material"], "thickness_mm": price["thickness_mm"],
              "width_mm": 300, "height_mm": 120, "language": "en",
              "evidence": [*schedule["evidence"], *drawing["evidence"]]}
    sequence = [
        call("set_reading_plan", plan="Schedule provides quantities; compare drawing and scope notes"),
        *[call("read_source", file_id=fid, page=1) for fid in (excel_id, pdf_id, image_id)],
        call("save_observations", observations=[
            {"line": schedule, "reason": "Schedule order row"}, {"line": drawing, "reason": "Related drawing"}]),
        call("reconcile_observations", action="merge", observation_ids=["o000001", "o000002"],
             merged_line=merged, reason="Compared location and drawing, same sign despite different titles"),
        *coverage_calls([excel_id, pdf_id, image_id]),
        call("search_prices", query=price["product"]),
    ]
    seen = []
    monkeypatch.setattr(deepseek, "get_key", lambda: "test-key-not-real")

    def turn(messages, tools):
        seen.append(copy.deepcopy(messages))
        if sequence:
            return sequence.pop(0), {"total_tokens": 10}
        context = json.loads(messages[1]["content"])
        identifier = context["requirements"][0]["id"]
        name = "finish_quote" if context["quote_progress"]["priced_count"] else "select_price"
        return (call(name, **({"line_id": identifier, "price_id": price["id"], "revision": price["revision"],
                             "reason": "Compared material and dimensions"} if name == "select_price" else {})),
                {"total_tokens": 10})

    monkeypatch.setattr(deepseek, "agent_turn", turn)
    result = start(client, job).json()
    assert result["agent_runs"][0]["status"] == "succeeded", result["agent_runs"][0]["message"]
    assert len(result["requirements"]) == 1
    assert result["requirements"][0]["quantity"] == 2
    assert {e["file_id"] for e in result["requirements"][0]["evidence"]} == {excel_id, pdf_id}
    assert result["quotes"][0]["payload"]["total"] == "21.60"
    assert result["quotes"][0]["status"] == "draft"
    assert result["extraction"]["source_agent"]["reviewed_pages"] == 3
    assert not result["quotes"][0]["payload"]["source_review"]["required"]
    assert client.get(f'/api/quotes/{result["quotes"][0]["id"]}/export').status_code == 409


def test_background_explicit_resume_grants_budget_and_reuses_read_sources(
        client, monkeypatch, redis_worker, test_db):
    job = create_job(client)
    fid = job["files"][0]["id"]
    seen = provider(monkeypatch, [
        call("read_source", file_id=fid, page=1),
        call("save_observations", observations=[{"line": line(fid), "reason": "Read sign"}]),
        *review_calls([fid])])
    monkeypatch.setattr(settings(), "agent_background", True)
    monkeypatch.setattr(settings(), "source_agent_max_tokens", 10)
    started = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": job["revision"], "allow_external_processing": True}).json()
    agent_worker.run_pending_once(threading.Event())
    paused = client.get(f'/api/jobs/{job["id"]}').json()
    assert paused["agent_runs"][0]["status"] == "waiting"
    assert paused["extraction"]["incomplete"]
    monkeypatch.setattr(settings(), "source_agent_max_tokens", 1000)
    resumed = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": paused["revision"], "allow_external_processing": True, "intent": "resume"})
    assert resumed.status_code == 200, resumed.text
    for _ in range(5):
        agent_worker.run_pending_once(threading.Event())
        finished = client.get(f'/api/jobs/{job["id"]}').json()
        if finished["status"] != "processing":
            break
    assert finished["agent_runs"][0]["status"] == "succeeded", finished["agent_runs"][0]["message"]
    assert len(finished["requirements"]) == 1
    assert len(seen) == 6
    with test_db() as db:
        runs = list(db.scalars(select(AgentRun).order_by(AgentRun.id)))
        assert sum(r.usage.get("total_tokens", 0) for r in runs) == 60
    assert started["agent_runs"][0]["id"] != finished["agent_runs"][0]["id"]


def test_reextract_preserves_human_price_and_id_through_final_publication(client, monkeypatch):
    job = create_job(client)
    fid = job["files"][0]["id"]
    raw = line_input(manual_unit_price="99.00", manual_price_note="Approved manually")
    raw["evidence"] = line(fid)["evidence"]
    job = save_lines(client, job, [raw])
    provider(monkeypatch, [call("read_source", file_id=fid, page=1), *review_calls([fid])])
    result = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": job["revision"], "allow_external_processing": True, "replace_existing": True}).json()
    assert result["requirements"][0]["id"] == "line1"
    assert result["requirements"][0]["confirmed"]
    assert result["requirements"][0]["manual_unit_price"] == "99.00"


def test_partial_quote_resume_reviews_remaining_pages_and_preserves_selected_price(
        client, monkeypatch, price_input):
    from test_agent_loop import seed
    price = seed(client, price_input)
    job = create_job(client, filename="drawings.pdf", data=document_bytes(pages=2, text=True))
    fid = job["files"][0]["id"]
    raw = line(fid, price["product"], material=price["material"],
               thickness_mm=price["thickness_mm"], width_mm=300, height_mm=120, language="en")
    sequence = [
        call("read_source", file_id=fid, page=1),
        call("save_observations", observations=[{"line": raw, "reason": "Read order drawing"}]),
        *review_calls([fid]),
        call("search_prices", query=price["product"]),
    ]
    monkeypatch.setattr(deepseek, "get_key", lambda: "test-key-not-real")

    def turn(messages, tools):
        if sequence:
            return sequence.pop(0), {"total_tokens": 10}
        context = json.loads(messages[1]["content"])
        current = context["requirements"][0]
        return (call("finish_quote") if current.get("selected_price_id") else
                call("select_price", line_id=current["id"], price_id=price["id"],
                     revision=price["revision"], reason="Compared actual material and size")), {"total_tokens": 10}

    monkeypatch.setattr(deepseek, "agent_turn", turn)
    partial = start(client, job).json()
    assert partial["agent_runs"][0]["status"] == "waiting", partial["agent_runs"][0]["message"]
    assert partial["quotes"][0]["payload"]["known_subtotal"] == "21.60"
    assert partial["quotes"][0]["payload"]["total"] is None
    assert partial["extraction"]["source_agent"]["unread_pages"] == [{"file_id": fid, "page": 2}]
    identifier = partial["requirements"][0]["id"]
    signature = partial["extraction"]["source_agent"]["input_signature"]
    sequence.extend([
        call("read_source", file_id=fid, page=2, purpose="review"),
        call("record_source_reviews", pages=[{"file_id": fid, "page": 2, "role": "note",
             "status": "reviewed", "reason": "Reference notes do not add ordered products"}]),
        call("finish_source_review", summary="All original pages compared"),
    ])
    finished = start(client, partial, intent="resume").json()
    assert finished["agent_runs"][0]["status"] == "succeeded", finished["agent_runs"][0]["message"]
    assert finished["requirements"][0]["id"] == identifier
    assert finished["requirements"][0]["selected_price_id"] == price["id"]
    audit = finished["extraction"]["source_agent"]
    assert audit["input_signature"] == signature
    assert audit["reviewed_pages"] == 2 and not audit["unreviewed_pages"]
    assert finished["quotes"][0]["payload"]["total"] == "21.60"


@pytest.mark.parametrize("durable", [False, True])
def test_pricing_original_pixels_survive_context_compaction_and_worker_slice(
        client, monkeypatch, redis_worker, durable):
    from app import agent_loop
    job = create_job(client)
    fid = job["files"][0]["id"]
    raw = line_input()
    raw["evidence"] = line(fid)["evidence"]
    job = save_lines(client, job, [raw])
    first = call("view_source_region", file_id=fid, page=1, region=[0, 0, 1, .5])
    second = call("view_source_region", file_id=fid, page=1, region=[0, .5, 1, 1])
    second["tool_calls"][0]["id"] = "second-region"
    first["tool_calls"].extend(second["tool_calls"])
    seen = provider(monkeypatch, [first, call("ask_user", question="Please confirm included production scope")])
    monkeypatch.setattr(settings(), "agent_background", durable)
    monkeypatch.setattr(settings(), "agent_slice_rounds", 1)
    monkeypatch.setattr(agent_loop, "MAX_CONTEXT_CHARS", 1)
    result = start(client, job, intent="resume").json()
    if durable:
        for _ in range(4):
            agent_worker.run_pending_once(threading.Event())
            result = client.get(f'/api/jobs/{job["id"]}').json()
            if result["status"] != "processing":
                break
    assert len(seen) == 2, result["agent_runs"][0]["message"]
    images = [part for message in seen[1] if isinstance(message.get("content"), list)
              for part in message["content"] if part.get("type") == "image_url"]
    assert len(images) == 2
    run = result["agent_runs"][0]
    assert run["status"] == "waiting"
    assert "_source_images" not in json.dumps(run) and "data:image" not in json.dumps(run)
    assert not run["usage"].get("source_view_pending")


def test_failed_replay_is_reported_to_ai_and_does_not_strand_resume(monkeypatch):
    file = file_fixture()
    cache, _, _ = checkpoint()
    session = SourceSession([file], "", [], "", cache)
    session.state["pending_sources"] = [{"tool": "read_source", "args": {"file_id": 99, "page": 1}}]
    cache.source_agent = session.state
    seen = provider(monkeypatch, [call("read_source", file_id=1, page=1), *review_calls([1])])
    result, _, _ = deepseek.extract([file], "", [], "", _checkpoint=cache)
    assert "回读失败" in seen[0][-1]["content"][0]["text"]
    assert result.source_agent["reviewed_pages"] == 1


def test_ai_observations_are_never_filtered_by_legacy_heading_heuristics():
    from app.non_product_rows import partition_requirements
    job = SimpleNamespace(extraction={"source_agent": {"mode": "ai"}},
                          requirements=[line(product="AI checked actual product")])
    kept, excluded = partition_requirements(None, job)
    assert kept == job.requirements and not excluded


def test_partial_pixels_keep_page_unread_even_after_ai_finishes_review():
    file = file_fixture()
    session = SourceSession([file], "", [], "", None)
    execute(session, "begin_source_review", summary="Read original top half")
    execute(session, "inspect_source_region", file_id=1, page=1, region=[0, 0, 1, .5], purpose="review")
    execute(session, "record_source_reviews", pages=[{
        "file_id": 1, "page": 1, "role": "drawing", "status": "reviewed", "reason": "Top half checked"}])
    execute(session, "finish_source_review", summary="Deliver usable scope")
    assert session.audit()["read_pages"] == 1
    assert source_review(session.result().model_dump(mode="json"))["unread_pages"] == 1


def test_batched_source_reads_deliver_both_sources_without_orphaned_tool_replies(monkeypatch):
    first = call("read_source", file_id=1, page=1)
    second = call("read_source", file_id=2, page=1)
    second["tool_calls"][0]["id"] = "second-source"
    first["tool_calls"].extend(second["tool_calls"])
    seen = provider(monkeypatch, [first, *review_calls([1, 2])])
    deepseek.extract([file_fixture(), file_fixture(2)], "", [], "")
    images = [part for message in seen[1] if isinstance(message.get("content"), list)
              for part in message["content"] if part.get("type") == "image_url"]
    assert len(images) == 2
    assert all(message["role"] != "tool" for message in seen[1])
    assert len(json.loads(seen[1][-1]["content"])["recent_tool_results"]) == 2


def test_synchronous_budget_pause_is_resumable_instead_of_failed(client, monkeypatch):
    job = create_job(client)
    fid = job["files"][0]["id"]
    monkeypatch.setattr(settings(), "source_agent_max_tokens", 10)
    provider(monkeypatch, [call("read_source", file_id=fid, page=1)])
    result = start(client, job).json()
    assert result["agent_runs"][0]["status"] == "waiting"
    assert result["extraction"]["incomplete"]
    assert "预算" in result["agent_runs"][0]["message"]


def test_script_observations_bypass_legacy_semantics_but_are_never_auto_adopted(monkeypatch):
    from app import file_sandbox, file_workspace
    from test_file_workspace import source_file
    file = source_file()
    session = SourceSession([file], "", [], "", None)
    output = {"observations": [{"name": "Inspect possible cover", "quantity": 99}]}
    monkeypatch.setattr(file_sandbox, "run_python", lambda *args: {"ok": True, "output": output})
    def reject_legacy(*args, **kwargs):
        raise AssertionError("AI observation mode must not invoke semantic hard validation")
    monkeypatch.setattr(file_workspace, "validate_output", reject_legacy)
    result = execute(session, "run_file_script", file_ids=[file.id],
                     script="print('inspect original')", purpose="Locate relevant original rows")
    assert result["status"] == "observations_only"
    assert not session.result().lines
    artifact = execute(session, "read_workspace_artifact", artifact_id=result["artifact_id"])
    assert json.loads(artifact["text"])["execution"]["output"] == output


def test_page_notes_survive_resume_without_claiming_omission_review_or_replaying_old_pixels(monkeypatch):
    file = file_fixture()
    seen = provider(monkeypatch, [
        call("read_source", file_id=1, page=1),
        call("record_source_notes", pages=[{"file_id": 1, "page": 1, "role": "cover",
             "summary": "Cover locates drawings; there are no ordered products here",
             "next_action": "Compare later order schedule"}]),
        call("list_source_notes"),
        *review_calls([1]),
    ])
    monkeypatch.setattr(settings(), "agent_slice_rounds", 1)
    saved = {}
    for _ in range(12):
        cache = ExtractionCheckpoint("notes", saved, lambda state, usage: saved.update(copy.deepcopy(state)), lambda: None)
        try:
            result, _, _ = deepseek.extract([file], "", [], "", _checkpoint=cache)
        except ExtractionYield:
            if saved["source_agent"]["source_notes"]:
                assert saved["source_agent"]["reviews"] == {} or saved["source_agent"]["phase"] == "review"
            continue
        break
    assert result.source_agent["source_notes"][0]["role"] == "cover"
    context = json.loads(seen[2][1]["content"])
    assert context["recent_source_notes"][0]["next_action"] == "Compare later order schedule"
    assert not any(isinstance(message.get("content"), list) for message in seen[2])
    assert len(result.lines) == 0 and result.source_agent["reviewed_pages"] == 1


def test_saving_observations_and_notes_is_atomic_and_releases_source_pixels():
    from app.source_agent_context import SourceContext, context_size
    session = SourceSession([file_fixture()], "", [], "", None)
    context = SourceContext(session)
    session.state["pending_sources"] = []
    args = SPECS["read_source"][0](file_id=1, page=1)
    result = session.execute("read_source", args)
    context.capture("read_source", args, result, list(session.images))
    assert context_size(context.messages())["images"] == 1
    with pytest.raises(ValueError, match="不属于"):
        execute(session, "save_observations",
            observations=[{"line": line(), "reason": "Visible product"}],
            source_notes=[{"file_id": 999, "page": 1, "role": "drawing", "summary": "Invalid source"}])
    assert not session.state["observations"]
    execute(session, "save_observations",
        observations=[{"line": line(quantity=None), "reason": "Save dimensions now; quantity awaits schedule"}],
        source_notes=[{"file_id": 1, "page": 1, "role": "drawing",
                       "summary": "Saved product; order quantity awaits the schedule"}])
    assert len(session.state["observations"]) == 1
    assert session.result().lines[0].quantity is None
    assert context_size(context.messages())["images"] == 0
    assert session.audit()["reviewed_pages"] == 0


def test_long_reading_uses_bounded_context_and_keeps_every_original_for_the_next_turn(monkeypatch):
    from app import source_agent
    from app.source_agent_context import context_size
    monkeypatch.setattr(source_agent, "pdf_page_text",
        lambda file, page: f"PAGE_DATA_{page:04d} " + "x" * 5600)
    file = file_fixture(filename="long.pdf", data=document_bytes(pages=18))
    sequence = [call("read_source", file_id=1, page=page) for page in range(1, 19)]
    sequence.extend(review_calls([1]))
    seen = provider(monkeypatch, sequence)
    result, _, _ = deepseek.extract([file], "", [], "")
    assert result.source_agent["read_pages"] == 18
    assert max(context_size(messages)["text_chars"] for messages in seen) < 24000
    for page in range(1, 19):
        raw = [json.loads(message["content"][0]["text"])["working_source"]
               for message in seen[page] if isinstance(message.get("content"), list)]
        assert any(source.get("page") == page and f"PAGE_DATA_{page:04d}" in source.get("text", "")
                   for source in raw)
        assert context_size(seen[page])["images"] <= 2
    assert "PAGE_DATA_0001" not in json.dumps(seen[18])
    assert "PAGE_DATA_0018" in json.dumps(seen[18])


def test_source_lists_are_paginated_and_schema_does_not_duplicate_complete_line_contract():
    from app.source_agent_tools import tools
    session = SourceSession([file_fixture(filename="long.pdf", data=document_bytes(pages=40))], "", [], "", None)
    result = execute(session, "list_sources", limit=10)
    assert len(result["pages"]) == 10 and result["next_offset"] == 10 and result["total"] == 40
    second = execute(session, "list_sources", offset=10, limit=10)
    assert second["pages"][0]["page"] == 11
    definitions = tools()
    # Preserve the complete save schema, without advertising a second copy for merge.
    save_schema = next(t["function"]["parameters"] for t in definitions if t["function"]["name"] == "save_observations")
    assert "components" in save_schema["$defs"]["ExtractedLine"]["properties"]
    assert len(json.dumps(definitions, ensure_ascii=False)) < 23000


def test_legacy_checkpoint_keeps_all_fresh_sources_across_failed_response(monkeypatch):
    from app.source_agent_context import context_size
    files = [file_fixture(file_id) for file_id in range(1, 5)]
    cache, snapshots, _ = checkpoint()
    session = SourceSession(files, "", [], "", cache)
    session.state["pending_sources"] = []
    for file in files:
        args = SPECS["read_source"][0](file_id=file.id, page=1)
        session.execute("read_source", args)
        session.state["pending_sources"].append({"tool": "read_source", "args": args.model_dump(mode="json")})
    # Older paid checkpoints contain pending originals, without the new memory fields.
    session.state.pop("source_notes")
    cache.source_agent = session.state
    failed_requests = []

    def fail(messages, definitions):
        failed_requests.append(copy.deepcopy(messages))
        raise deepseek.AgentResponseFailure("truncated", {"finish_reason": "length"}, {"total_tokens": 17})

    monkeypatch.setattr(deepseek, "agent_turn", fail)
    with pytest.raises(deepseek.ModelFailure):
        deepseek.extract(files, "", [], "", _checkpoint=cache)
    assert context_size(failed_requests[0])["images"] == 4
    saved = snapshots[-1]
    assert len(saved["source_agent"]["pending_sources"]) == 4
    assert saved["source_agent"]["total_tokens"] == 51
    assert "data:image" not in json.dumps(saved)

    notes = [{"file_id": file.id, "page": 1, "role": "cover",
              "summary": "Cover saved before continuing"} for file in files]
    seen = provider(monkeypatch, [call("record_source_notes", pages=notes),
                                *review_calls([file.id for file in files])])
    resumed, _, _ = checkpoint(saved)
    resumed.allow_yield = False
    result, _, _ = deepseek.extract(files, "", [], "", _checkpoint=resumed)
    assert context_size(seen[0])["images"] == 4
    assert not any(isinstance(message.get("content"), list) for message in seen[1])
    assert len(result.source_agent["source_notes"]) == 4
    assert result.source_agent["reviewed_pages"] == 4


def test_invalid_json_retries_single_item_without_executing_partial_batch(monkeypatch):
    from test_agent_response_recovery import envelope
    good = call("save_observations", observations=[{"line": line(product="Must not be saved"),
                                                   "reason": "Malformed batch"}])
    bad = call("save_observations")
    bad["tool_calls"][0]["id"] = "invalid-second-call"
    bad["tool_calls"][0]["function"]["arguments"] = '{"observations":'
    good["tool_calls"].extend(bad["tool_calls"])
    sequence = iter([
        call("read_source", file_id=1, page=1), good,
        call("save_observations", observations=[{"line": line(), "reason": "Save one product"}]),
        *review_calls([1]),
    ])
    requests = []

    def respond(messages, definitions):
        requests.append((copy.deepcopy(messages), copy.deepcopy(definitions)))
        return deepseek.parse_agent_response(envelope(next(sequence)))

    monkeypatch.setattr(deepseek, "agent_turn", respond)
    cache, snapshots, charges = checkpoint()
    cache.allow_yield = False
    result, attempts, usage = deepseek.extract([file_fixture()], "", [], "", _checkpoint=cache)
    assert [row.product for row in result.lines] == ["Sign"]
    assert result.source_agent["observation_ids"] == ["o000001"]
    assert attempts == len(requests) == 7 and usage["total_tokens"] == 119
    assert sum(item.get("total_tokens", 0) for item in charges) == 119
    recovering = json.loads(requests[2][0][1]["content"])["response_recovery"]
    assert recovering["active"] and recovering["executed_calls"] == 0
    schema = next(tool["function"]["parameters"] for tool in requests[2][1]
                  if tool["function"]["name"] == "save_observations")
    assert schema["properties"]["observations"]["maxItems"] == 1
    assert any(s["source_agent"]["response_recovery"].get("active") for s in snapshots)
    assert not snapshots[-1]["source_agent"]["response_recovery"]["active"]
    assert "Must not be saved" not in json.dumps(snapshots)


def test_protocol_failure_limit_survives_checkpoint_and_does_not_retry_forever(monkeypatch):
    def failed(*args):
        raise deepseek.AgentResponseFailure("invalid_arguments_json", {}, {"total_tokens": 17})
    monkeypatch.setattr(deepseek, "agent_turn", failed)
    cache, snapshots, charges = checkpoint()
    cache.allow_yield = False
    file = file_fixture()
    with pytest.raises(deepseek.ModelFailure) as first:
        deepseek.extract([file], "", [], "", _checkpoint=cache)
    assert first.value.diagnostics["retryable"] is True
    resumed, next_snapshots, next_charges = checkpoint(snapshots[-1])
    resumed.allow_yield = False
    with pytest.raises(deepseek.ModelFailure) as second:
        deepseek.extract([file], "", [], "", _checkpoint=resumed)
    assert second.value.diagnostics["retryable"] is False
    assert second.value.diagnostics["stage"] == "source_agent_response"
    assert next_snapshots[-1]["source_agent"]["response_recovery"]["consecutive_failures"] == 6
    assert sum(item.get("total_tokens", 0) for item in [*charges, *next_charges]) == 102


def test_schema_errors_return_field_paths_without_echoing_private_input(monkeypatch):
    sequence = iter([
        call("read_source", file_id=1, page=1),
        call("save_observations", observations=[{
            "line": line(material=None, customer_secret="DO_NOT_LOG"), "reason": "Invalid fields"}]),
        *review_calls([1]),
    ])
    seen = []

    def respond(messages, definitions):
        seen.append(copy.deepcopy(messages))
        return next(sequence), {"total_tokens": 1}

    monkeypatch.setattr(deepseek, "agent_turn", respond)
    result, _, _ = deepseek.extract([file_fixture()], "", [], "")
    latest = json.loads(seen[2][-1]["content"])["recent_tool_results"][0]["result"]
    assert "observations.0.line.material" in [item["path"] for item in latest["fields"]]
    assert json.loads(seen[2][1]["content"])["response_recovery"]["active"]
    assert "DO_NOT_LOG" not in json.dumps(latest)
    assert not result.lines


def test_background_real_saves_reset_old_network_retry_counts(
        client, monkeypatch, redis_worker, test_db):
    job = create_job(client, filename="order.pdf", data=document_bytes(pages=6))
    fid = job["files"][0]["id"]
    sequence = []
    for page in range(1, 7):
        sequence.extend([
            call("read_source", file_id=fid, page=page),
            deepseek.ModelFailure("Temporary network outage", diagnostics={"retryable": True}),
            call("save_observations", observations=[{"line": line(fid, f"Sign {page}", evidence=[
                {"file_id": fid, "page": page, "text": f"Sign {page}", "source": "image"}]),
                "reason": "Actual product saved after retry"}]),
        ])
    sequence.extend([
        call("begin_source_review", summary="Compare saved products"),
        *[call("read_source", file_id=fid, page=page, purpose="review") for page in range(1, 7)],
        call("record_source_reviews", pages=[{"file_id": fid, "page": page, "role": "drawing",
             "status": "reviewed", "reason": "Compared saved product"} for page in range(1, 7)]),
        call("finish_source_review", summary="Products saved and compared"),
    ])
    responses = iter(sequence)
    monkeypatch.setattr(deepseek, "get_key", lambda: "test-key-not-real")

    def respond(*args):
        response = next(responses)
        if isinstance(response, Exception):
            raise response
        return response, {"total_tokens": 1}

    monkeypatch.setattr(deepseek, "agent_turn", respond)
    monkeypatch.setattr(settings(), "agent_background", True)
    monkeypatch.setattr(settings(), "agent_retry_seconds", 0)
    monkeypatch.setattr(settings(), "agent_slice_rounds", 3)
    started = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": job["revision"], "allow_external_processing": True})
    assert started.status_code == 200, started.text
    for _ in range(35):
        agent_worker.run_pending_once(threading.Event())
        current = client.get(f'/api/jobs/{job["id"]}').json()
        if current["status"] != "processing":
            break
    run = current["agent_runs"][0]
    assert run["status"] == "succeeded", run["message"]
    assert len(current["requirements"]) == 6
    assert run["usage"]["retry_count"] == 0
    assert len([step for step in run["steps"] if step["tool"] == "extract_requirements_failed"]) == 6
    with test_db() as db:
        assert len(list(db.scalars(select(AgentRun).where(AgentRun.job_id == job["id"])))) == 1


def test_resuming_legacy_failed_run_starts_with_recovery_guidance(
        client, monkeypatch, redis_worker, test_db):
    job = create_job(client)
    fid = job["files"][0]["id"]
    monkeypatch.setattr(settings(), "agent_background", True)
    monkeypatch.setattr(settings(), "agent_slice_rounds", 1)
    monkeypatch.setattr(deepseek, "get_key", lambda: "test-key-not-real")
    provider(monkeypatch, [call("read_source", file_id=fid, page=1)])
    started = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": job["revision"], "allow_external_processing": True})
    assert started.status_code == 200, started.text
    agent_worker.run_pending_once(threading.Event())
    stopped = client.post(f'/api/jobs/{job["id"]}/agent/stop',
                          json={"revision": 1, "reason": "Simulate an older failed run"})
    assert stopped.status_code == 200, stopped.text
    with test_db.begin() as db:
        old = db.scalar(select(AgentRun).where(AgentRun.job_id == job["id"]))
        old_usage = copy.deepcopy(old.usage)
        old_usage["extraction_checkpoint"]["source_agent"].pop("response_recovery")
        old.usage = old_usage
        old.steps = [*old.steps, {"tool": "source_agent.response_failed", "arguments": {},
                     "result": {"diagnostics": {"reason": "invalid_arguments_json"}}}]
    seen = provider(monkeypatch, [
        call("record_source_notes", pages=[{"file_id": fid, "page": 1, "role": "cover",
                                           "summary": "Original cover checked"}]),
        *review_calls([fid]),
    ])
    resumed = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": stopped.json()["revision"], "allow_external_processing": True, "intent": "resume"})
    assert resumed.status_code == 200, resumed.text
    for _ in range(10):
        agent_worker.run_pending_once(threading.Event())
        current = client.get(f'/api/jobs/{job["id"]}').json()
        if current["status"] != "processing":
            break
    assert current["agent_runs"][0]["status"] == "succeeded", current["agent_runs"][0]["message"]
    recovery = json.loads(seen[0][1]["content"])["response_recovery"]
    assert recovery["active"] and recovery["diagnostics"]["reason"] == "invalid_arguments_json"
    assert not current["requirements"]


def test_background_saves_as_it_reads_and_continues_past_previous_cap_without_manual_resume(
        client, monkeypatch, redis_worker, test_db):
    job = create_job(client, filename="order.pdf", data=document_bytes(pages=12))
    fid = job["files"][0]["id"]
    sequence = []
    for page in range(1, 13):
        observed = line(fid, f"Sign {page}", evidence=[
            {"file_id": fid, "page": page, "text": f"Sign {page}", "source": "image"}])
        sequence.extend([
            call("read_source", file_id=fid, page=page),
            call("save_observations", observations=[{"line": observed, "reason": "Visible ordered product"}],
                 source_notes=[{"file_id": fid, "page": page, "role": "drawing",
                                "summary": "Saved product; compare during omission review"}]),
        ])
    sequence.extend([
        call("record_source_reviews", pages=[{"file_id": fid, "page": page, "role": "drawing",
             "status": "reviewed", "reason": "Initial read covers saved products and page layout"}
             for page in range(1, 13)]),
        call("finish_source_review", summary="All products compared"),
    ])
    monkeypatch.setattr(deepseek, "get_key", lambda: "test-key-not-real")
    calls = iter(sequence)
    monkeypatch.setattr(deepseek, "agent_turn", lambda *args: (next(calls), {"total_tokens": 12000}))
    monkeypatch.setattr(settings(), "agent_background", True)
    monkeypatch.setattr(settings(), "agent_slice_rounds", 3)
    response = client.post(f'/api/jobs/{job["id"]}/extract', json={
        "revision": job["revision"], "allow_external_processing": True})
    assert response.status_code == 200, response.text
    partial_counts = []
    for _ in range(20):
        agent_worker.run_pending_once(threading.Event())
        current = client.get(f'/api/jobs/{job["id"]}').json()
        partial_counts.append(len(current["requirements"]))
        if current["status"] != "processing":
            break
    assert current["agent_runs"][0]["status"] == "succeeded", current["agent_runs"][0]["message"]
    assert len(current["requirements"]) == 12
    assert any(0 < count < 12 for count in partial_counts)
    audit = current["extraction"]["source_agent"]
    assert audit["reviewed_pages"] == 12 and len(audit["source_notes"]) == 12
    with test_db() as db:
        runs = list(db.scalars(select(AgentRun).where(AgentRun.job_id == job["id"])))
        assert len(runs) == 1 and runs[0].usage["total_tokens"] == 312000
        assert len([step for step in runs[0].steps if step["tool"] == "source_agent.read_source"]) == 12
        assert not any(step["tool"] == "source_agent.budget_pause" for step in runs[0].steps)
        turn = next(step["result"] for step in runs[0].steps if step["tool"] == "source_agent.turn")
        assert turn["text_chars"] and turn["tool_schema_chars"]
        assert "data:image" not in json.dumps(runs[0].usage)
