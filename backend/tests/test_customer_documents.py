import base64
import gzip
import hashlib
import io
import json
import os
import struct
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PIL import Image

from app import file_sandbox, file_workspace
from app.customer_excel import MEDIA_TYPE, page_source
from app.customer_files import FileRejected, prepare_file, storage_path
from app.customer_documents import legacy_has_objects
from app.excel_manifest import ensure_preview
from app.file_workspace import FileScriptInput
from app.native_excel import extract_native
from app.source_agent import SourceSession
from app.source_agent_tools import ReadSource, RecordReviews
from app.source_reconciliation import source_review


def word_bytes(body=None, image=True):
    body = body or (
        '<w:p><w:r><w:t>项目说明</w:t></w:r></w:p>'
        '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>序号</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t>产品名称</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t>数量</w:t></w:r></w:p></w:tc></w:tr>'
        '<w:tr><w:tc><w:p><w:r><w:t>1</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t>标识牌</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t>2</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
    if image:
        body += '<w:p><w:r><w:drawing><a:blip r:embed="rId1"/></w:drawing></w:r></w:p>'
    document = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<w:body>{body}</w:body></w:document>')
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        archive.writestr("word/document.xml", document)
        archive.writestr("word/_rels/document.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Target="media/image1.png"/></Relationships>')
        if image:
            pixels = io.BytesIO()
            Image.new("RGB", (64, 32), "#89ab12").save(pixels, "PNG")
            archive.writestr("word/media/image1.png", pixels.getvalue())
    return stream.getvalue()


@pytest.mark.parametrize("extension,encoding,delimiter", [
    ("csv", "utf-8-sig", ","), ("csv", "gb18030", ";"), ("tsv", "utf-16", "\t"),
])
def test_text_tables_keep_exact_values_source_records_and_original(extension, encoding, delimiter):
    rows = [delimiter.join(["序号", "产品名称", "数量", "单位"]),
            delimiter.join(["001", "标识牌", "2", "套"])]
    data = "\r\n".join(rows).encode(encoding)
    file = SimpleNamespace(id=1, **prepare_file(data, f"source.{extension}"))
    directory = storage_path(file.storage_key)
    assert (directory / "original").read_bytes() == data
    assert file.sha256 == hashlib.sha256(data).hexdigest()
    source = page_source(directory, 1)
    assert source["rows"][1]["cells"][0]["value"] == "001"
    assert source["rows"][1]["cells"][1]["origin"]["record"] == 2
    result, _, _ = extract_native([file], "", [], "", None)
    assert len(result.lines) == 1 and result.lines[0].quantity == 2
    assert result.lines[0].product == "标识牌"
    assert ensure_preview(directory, 1).is_file()


def test_csv_keeps_multiline_quoted_fields_and_never_executes_formula_text():
    data = '序号,产品名称,数量,说明\n1,"入口,\n标识",2,=1+2'.encode()
    file = prepare_file(data, "quoted.csv")
    source = page_source(storage_path(file["storage_key"]), 1)
    values = [cell["value"] for cell in source["rows"][1]["cells"]]
    assert values == ["1", "入口,\n标识", "2", "=1+2"]
    assert "formula" not in source["rows"][1]["cells"][-1]


def test_empty_first_column_does_not_overwrite_prior_source_rows():
    file = prepare_file(b",Product,Quantity\n,Sign,2\n,Arrow,3", "empty-first-column.csv")
    source = page_source(storage_path(file["storage_key"]), 1)
    assert [next(cell["value"] for cell in row["cells"] if cell["cell"].startswith("B"))
            for row in source["rows"]] == ["Product", "Sign", "Arrow"]
    assert source["rows"][2]["cells"][0]["origin"]["record"] == 3


def test_word_preserves_paragraphs_table_order_images_and_original_locations():
    data = word_bytes()
    file = prepare_file(data, "specification.docx")
    directory = storage_path(file["storage_key"])
    source = page_source(directory, 1)
    assert file["media_type"] == MEDIA_TYPE
    assert source["rows"][0]["cells"][0]["value"] == "项目说明"
    cell = source["rows"][2]["cells"][1]
    assert cell["value"] == "标识牌" and cell["origin"]["part"] == "word/document.xml"
    assert cell["origin"]["table_row"] == 2
    assert source["images"][0]["cell"] == "A4"
    assert (directory / "original").read_bytes() == data
    assert "不是Word打印页码" in source["notice"]
    assert ensure_preview(directory, 1).is_file()


def test_plain_text_has_traceable_lines():
    file = prepare_file("计价规则\n加工费单独计算".encode(), "rules.txt")
    source = page_source(storage_path(file["storage_key"]), 1)
    assert source["rows"][1]["cells"][0]["value"] == "加工费单独计算"
    assert source["rows"][1]["cells"][0]["origin"]["record"] == 2


def test_missing_word_image_remains_an_unresolved_source_even_after_ai_review():
    data = word_bytes(image=False,
        body='<w:p><w:r><w:t>原图</w:t><w:drawing><a:blip r:embed="missing"/></w:drawing></w:r></w:p>')
    file = SimpleNamespace(id=1, **prepare_file(data, "missing-image.docx"))
    session = SourceSession([file], "", [], "", None)
    session.read(ReadSource(file_id=1, page=1))
    session.execute("record_source_reviews", RecordReviews(pages=[{
        "file_id": 1, "page": 1, "role": "drawing", "status": "reviewed", "reason": "AI review"}]))
    result = session.result().model_dump(mode="json")
    assert source_review(result)["required"]
    assert result["source_coverage"][0]["status"] == "unresolved"
    assert "未成功读取" in result["source_coverage"][0]["reason"]


def legacy_bytes():
    # A tiny BIFF workbook with Item/Product/Quantity and one Sign, quantity 2.
    packed = (
        "H4sIAAAAAAAC/+1YTWgTQRT+ZpM0P7TppqZCK5RQsGrbi3jx0m4VtHiwpPagiKBpstjQuJFlc6gHrdYcBcGT4qXQi5eqF3/Qg948CJV6EAQh0aMnQcFDm/XNy25Jq4cGNKjMt8ybt+/Nm3mZefMms29WEpXFh71VbMEoAqi5UbQ1yASVqP+ig/SuK1m/jlBxFf4pRCO0kG0hPOt4HZZrKNe7Cg0Pgi+JAh+pnMFFTBQtM9VCHGYfMkL6MEJU4C5J4uhhr7qYZpnuYHqfWz5nOsaSG0xHqG1FnMaKMTF40IviU1o/6+KQ/T5mm/cs2Y9uvJJRfOWmqLcN4ZCdzxT+TkVfsB1LoHUbNy3TzhQqSNICLuGbmwK++jv1RUrJWysXIPn3zfLwL+S3tCAwD/ccB3iZcupqoJ5Ip7IzZq5UMNcwBGrjl2OOeYGWO20Xc6WsQ/l2spSxnLwzR7qp/HkrJjM072h9047u4EhvJ5pDJ/MJjnedhlq792X1+HTaOMuSec7i9Vy/W7oHF1elBRnHWaMxlX4OMj/E9Br3uov5XqZJilKqB9LdHnN0gdtcZ+0AjXOA8dbY08DvJb78efJJX/mTsY/45fHqpeTyO2MR/XT25MhePgsYFsPizm2Jp4ZfCy8vfGDa81OOiGi657vrHWidWEeM2QTT+pucHbHxpnlzJa3FFuvLWoxlCZ5l2V6wtZyduk5ay15GtS484qkbazhRY1BQUFBQUFBQUFBoJYT3jz7gXTJC3o0h7H3XWadSU59J/lucQJEehy6mR2BRbWOuqfjZiZDw+xLbtPG/F0qcpNFtzGKa/ZhtOn7pxicaf8+2DfXft4WaHb/WjJ9/ePwfd2fmgwAWAAA=")
    return gzip.decompress(base64.b64decode(packed))


def test_legacy_xls_is_read_from_real_biff_content():
    file = prepare_file(legacy_bytes(), "legacy.xls")
    source = page_source(storage_path(file["storage_key"]), 1)
    assert source["rows"][1]["cells"][1]["value"] == "Sign"
    assert source["rows"][1]["cells"][2]["value"] == "2"
    assert source["rows"][1]["cells"][1]["origin"]["row"] == 2
    assert any("公式缓存值" in warning for warning in source["warnings"])


def test_legacy_xls_drawing_records_are_detected():
    plain = struct.pack("<HHH", 0x0203, 2, 0)
    objects = plain + struct.pack("<HH", 0x00EC, 0)
    assert not legacy_has_objects(SimpleNamespace(base=0, stream_len=len(plain), mem=plain))
    assert legacy_has_objects(SimpleNamespace(base=0, stream_len=len(objects), mem=objects))


@pytest.mark.parametrize("name,data", [
    ("bad.csv", b"\x00\x01\x02"), ("bad.docx", b"not a zip"),
    ("bad.xls", b"not an xls"), ("bad.docx", word_bytes('<!DOCTYPE evil>', image=False)),
])
def test_invalid_documents_are_rejected_without_leaving_uploads(name, data):
    with pytest.raises(FileRejected):
        prepare_file(data, name)


def test_converted_workspace_exposes_original_and_normalized_files(monkeypatch):
    data = b"Item,Product,Quantity\n1,Sign,2"
    file = SimpleNamespace(id=1, job_id=7, **prepare_file(data, "source.csv"))

    def execute(directory, check):
        manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
        entry = manifest["files"][0]
        assert entry["path"] == "/input/file-1.csv" and entry["media_type"] == "text/csv"
        assert (directory / "file-1.csv").read_bytes() == data
        assert entry["normalized_path"] == "/input/file-1-normalized.xlsx"
        workbook = load_workbook(directory / "file-1-normalized.xlsx")
        try:
            assert workbook.active["B2"].value == "Sign"
        finally:
            workbook.close()
        return {"ok": True, "output": {"observations": ["Sign"]}}

    monkeypatch.setattr(file_sandbox, "run_python", execute)
    result = file_workspace.run_script(7, [file], FileScriptInput(file_ids=[1], script="pass", purpose="Read"),
                                       observation_mode=True)
    assert result["status"] == "observations_only"
    assert not list((file_workspace.job_root(7) / result["artifact_id"] / "input").glob("file-*"))


@pytest.mark.skipif(os.environ.get("FIREFLY_SANDBOX_INTEGRATION") != "1", reason="Explicit Docker integration run")
@pytest.mark.parametrize("extension", ["csv", "docx", "xls"])
def test_real_container_reads_converted_original_and_normalized_sources(extension):
    data = {"csv": lambda: b"Item,Product,Quantity\n1,Sign,2",
            "docx": lambda: word_bytes(image=False), "xls": legacy_bytes}[extension]()
    file = SimpleNamespace(id=1, job_id=7, **prepare_file(data, f"source.{extension}"))
    script = '''
import hashlib, json
import xlrd
from openpyxl import load_workbook
m = json.load(open("/input/manifest.json", encoding="utf-8"))
f = m["files"][0]
assert hashlib.sha256(open(f["path"], "rb").read()).hexdigest() == f["sha256"]
w = load_workbook(f["normalized_path"], data_only=True)
values = [str(c.value) for row in w.active for c in row if c.value is not None]
assert len(values) >= 6
assert any(cell.get("origin") for p in m["source_pages"]
           for row in json.loads(p["text"])["rows"] for cell in row["cells"])
json.dump({"observations": values}, open("/work/result.json", "w", encoding="utf-8"))
'''
    result = file_workspace.run_script(7, [file], FileScriptInput(
        file_ids=[1], script=script, purpose="Verify converted source in real container"), observation_mode=True)
    assert result["status"] == "observations_only", result
    assert not list((file_workspace.job_root(7) / result["artifact_id"] / "input").glob("file-*"))


def test_normalized_documents_have_working_source_and_reparse_endpoints(client):
    response = client.post("/api/jobs", data={"title": "CSV order"},
                           files={"files": ("source.csv", "序号,产品名称,数量\n1,标识牌,2".encode(), "text/csv")})
    assert response.status_code == 201, response.text
    file = response.json()["files"][0]
    assert file["conversion"]["original_media_type"] == "text/csv"
    source = client.get(f'/api/customer-files/{file["id"]}/source/1')
    assert source.status_code == 200 and source.json()["rows"][1]["cells"][1]["origin"]["record"] == 2
    preview = client.get(f'/api/customer-files/{file["id"]}/pages/1')
    assert preview.status_code == 200, preview.text[:100]
    reparse = client.post(f'/api/customer-files/{file["id"]}/reparse-preview')
    assert reparse.status_code == 200, reparse.text
